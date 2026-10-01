import csv
import importlib
from io import BytesIO, StringIO

import pytest
from openpyxl import load_workbook
from starlette.testclient import TestClient

from ims.api.app import create_app
from tests.test_ict_workshop import small_case


@pytest.mark.parametrize("fallback", [False, True])
def test_ict_api_exact_exports_and_no_metadata_or_market_execution(monkeypatch, tmp_path, fallback) -> None:
    module = importlib.import_module("ims.api.app")
    if fallback:
        monkeypatch.setattr(module, "FastAPI", None)
    db = tmp_path / "not-created.sqlite"
    monkeypatch.setenv("IMS_METADATA_DB", str(db))
    client = TestClient(create_app(frontend_dist=tmp_path))
    doc = small_case()
    assert client.get("/api/ict/contract").json()["compliance_decision_enabled"] is False
    assert client.post("/api/ict/validate", json=doc).json()["ict_model_calculated"] is False
    response = client.post("/api/ict/calculate", json=doc)
    assert response.status_code == 200
    result = response.json()
    assert response.headers["etag"] == '"' + result["content_digest"] + '"'
    for extension in ("json", "csv", "xlsx"):
        exported = client.post(f"/api/ict/export.{extension}", json=doc, headers={"If-Match": response.headers["etag"]})
        assert exported.status_code == 200
        assert exported.headers["etag"] == response.headers["etag"]
        assert exported.headers["cache-control"] == "no-store"
        if extension == "json":
            assert exported.json() == result
        elif extension == "csv":
            rows = list(csv.DictReader(StringIO(exported.text)))
            assert len(rows) == 6
            assert rows[-1]["model_own_funds_proxy"] == "7297.1200"
            assert all(r["content_digest"] == result["content_digest"] for r in rows)
        else:
            book = load_workbook(BytesIO(exported.content), data_only=False)
            assert book["Variante Bilanz"].max_row == 4
            assert book["Herkunft"]["B2"].value == result["content_digest"]
            assert book["Variante Bilanz"]["H4"].value == "7297.1200"
    assert not db.exists()
    assert client.post("/api/ict/export.json", json=doc).status_code == 412
    doc["period_count"] = 2
    assert client.post("/api/ict/export.json", json=doc, headers={"If-Match": response.headers["etag"]}).status_code == 412
    invalid = client.post("/api/ict/calculate", content="{", headers={"Content-Type": "application/json"})
    assert invalid.status_code == 400 and invalid.json()["partial_result_returned"] is False
    assert client.post("/api/ict/calculate", content=" " * 256001).status_code == 400
