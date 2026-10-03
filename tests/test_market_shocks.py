"""AP7 causal processing/guarantee/accounting cases and fresh transport proofs."""
from __future__ import annotations

from copy import deepcopy
from decimal import Decimal
from io import BytesIO
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from starlette.testclient import TestClient

from ims.api.market import create_market_app
from ims.market.presets import workshop_case
from ims.market.runner import side_result
from ims.market.shock_contract import validate
from ims.market.shock_plan import ShockPlan
from ims.market.shock_presets import build_case
from ims.market.shock_process import CostLedger, Job, ProcessEngine
from ims.strategies.modern_bridge import ContractError

ROOT = Path(__file__).resolve().parents[1]


def process_bundle(*, rate="1", events=None, responses=None):
    assets = [{"asset_id": key, "control": control, "depends_on": dependencies, "label": key}
              for key, control, dependencies in (("primary", "US", []), ("independent", "EU", []), ("dependent", "EU", ["primary"]))]
    services = [{"service_id": key, "insurer_id": aid, "sector_id": "motor", "kind": "underwriting", "resource_id": "shared", "arrivals_per_period": 0}
                for key, aid in (("U1", 1), ("U2", 2))]
    return {"ict": {"assets": assets, "resources": [{"resource_id": "shared", "asset_id": "primary", "capacity_per_hour": rate,
             "cost_per_hour": "0", "owners": [{"insurer_id": 1, "sector_id": "motor", "weight": "1"}]}],
             "services": services, "holding_cost_per_job_hour": "0"}, "events": events or [], "responses": responses or []}


def outage(start="120", duration="24", intensity="1"):
    return {"event_id": "failure", "kind": "provider_outage", "start_hour": start, "duration_hours": duration,
            "intensity": intensity, "asset_ids": ["primary"], "insurer_ids": [1, 2], "channels": ["capacity"], "cost": "0", "assumption_note": "Handfall"}


def engine(bundle):
    return ProcessEngine(bundle, "variant", bundle["events"], bundle["responses"], CostLedger())


def test_half_open_subperiod_clock_preserves_shared_work():
    worker = engine(process_bundle(events=[outage("26", "2")]))
    done = worker.period(2, [Job(f"A{i:02}", "U1" if i % 2 else "U2", 1 if i % 2 else 2, "motor", "switch", 2) for i in range(40)], {1, 2})
    assert len(done) == 22
    assert worker.resource_rows[0]["capacity_work"] == "22.0000"
    assert worker.resource_rows[0]["processed_work"] == "22.0000"
    assert worker.resource_rows[0]["closing_queue"] == 18
    assert sum(row["completed"] for row in worker.process_rows) == 22
    assert worker.state("primary", Decimal(28))[0] == 1


def test_fractional_progress_survives_period_boundary():
    worker = engine(process_bundle(rate="0.05"))
    done = worker.period(1, [Job("A", "U1", 1, "motor", "switch", 1), Job("B", "U2", 2, "motor", "switch", 1)], {1, 2})
    assert [job.job_id for job in done] == ["A"]
    assert worker.queues["shared"][0].remaining_work == Decimal("0.8")
    done = worker.period(2, [], {1, 2})
    assert [job.job_id for job in done] == ["B"]
    assert done[0].completed_hour == Decimal(40)


def test_fallback_depends_on_entire_replacement_path():
    base = {"response_id": "Q", "kind": "fallback", "decision_period": 6, "lead_periods": 0, "duration_periods": 2,
            "asset_id": "primary", "replacement_asset_id": "independent", "value": "1"}
    for replacement, expected in (("independent", 2), ("dependent", 0)):
        response = {**base, "replacement_asset_id": replacement}
        worker = engine(process_bundle(events=[outage()], responses=[response]))
        done = worker.period(6, [Job("A", "U1", 1, "motor", "life_application", 6), Job("B", "U2", 2, "motor", "switch", 6)], {1, 2})
        assert len(done) == expected
    source = process_bundle(events=[outage()], responses=[base])
    source["ict"]["assets"][1]["control"] = "unknown"
    worker = engine(source)
    assert worker.period(6, [Job("A", "U1", 1, "motor", "switch", 6)], {1, 2}) == []


def test_overlapping_capacity_failures_do_not_add_and_costs_conserve():
    events = [outage(intensity="0.5"), {**outage(intensity="0.75"), "event_id": "second"}]
    worker = engine(process_bundle(events=events))
    assert worker.state("primary", Decimal(120))[0] == Decimal("0.25")
    costs = CostLedger()
    owners = [{"insurer_id": i, "sector_id": "motor", "weight": weight} for i, weight in enumerate(("0.3333", "0.3333", "0.3334"), 1)]
    costs.allocate("one", 1, Decimal(6), owners, "measure", "hand")
    assert [row["amount"] for row in costs.rows] == ["1.9998", "1.9998", "2.0004"]
    with pytest.raises(ContractError):
        costs.allocate("one", 1, Decimal(6), owners, "measure", "hand")
    tiny = CostLedger()
    tiny.allocate("tiny", 1, Decimal("0.0021"), [{"insurer_id": i, "sector_id": "motor", "weight": "0.025"} for i in range(40)], "operation", "rounding")
    assert sum(Decimal(row["amount"]) for row in tiny.rows) == Decimal("0.0021")
    assert all(Decimal(row["amount"]) >= 0 for row in tiny.rows)


def tiny_market(*, life=False):
    doc = workshop_case("market" if life else "switch", 10, 2)
    doc["assignments"]["variant"] = deepcopy(doc["assignments"]["baseline"])
    doc["measures"] = {"baseline": [], "variant": []}
    bundle = process_bundle(events=[outage()] if not life else [])
    bundle.update(model_input=doc, comparison="shock_with_response", activation=[],
                  responses={"baseline": [], "variant": []}, life={"pool_per_period": 0, "start_period": 6, "offers": []})
    if life:
        for aid in (1, 2):
            bundle["ict"]["services"].append({"service_id": f"UP{aid}", "insurer_id": aid, "sector_id": "property_liability", "kind": "underwriting", "resource_id": "shared", "arrivals_per_period": 0})
        bundle["ict"]["services"].append({"service_id": "life1", "insurer_id": 1, "sector_id": "life", "kind": "underwriting", "resource_id": "shared", "arrivals_per_period": 0})
        bundle["life"] = {"pool_per_period": 1, "start_period": 6, "offers": [{"insurer_id": 1, "premium": "3", "allocation": "2", "renewal_premium": "3", "renewal_allocation": "2", "term_periods": 2, "guarantee_rate": "0.001", "admission_per_period": 1}]}
    return bundle


def test_pending_switch_keeps_old_risk_and_premium_then_transfers_once():
    source = tiny_market()
    before = deepcopy(source)
    plan = ShockPlan(source, "baseline")
    result = side_result(source["model_input"], "baseline", {}, plan)
    decisions = {row["period"]: row for row in result["customer_decisions"]}
    assert decisions[6]["insurer_id"] == decisions[7]["insurer_id"] == 1
    assert decisions[6]["premium_income"] == "30.0000"
    assert decisions[6]["risk_loss"] == "40.0000"
    assert decisions[8]["insurer_id"] == 2 and decisions[8]["premium_income"] == "20.0000"
    assert plan.switch_rows[0]["request_period"] == 6 and plan.switch_rows[0]["effective_period"] == 8
    for row in result["vu_rows"]:
        assert Decimal(row["closing_assets"]) == Decimal(row["closing_liabilities"]) + Decimal(row["closing_equity"])
    assert source == before


def test_life_issue_waits_for_processing_and_existing_core_pays_guarantee():
    source = tiny_market(life=True)
    before = deepcopy(source)
    plan = ShockPlan(source, "baseline")
    result = side_result(source["model_input"], "baseline", {}, plan)
    assert plan.issues[0]["request_period"] == 6 and plan.issues[0]["issue_period"] == 7
    assert plan.issues[0]["maturity_period"] == 9
    rows = {row["period"]: row for row in result["vu_rows"] if row["insurer_id"] == 1 and row["sector_id"] == "life"}
    assert rows[6]["ap7_new_policies"] == 0 and rows[7]["ap7_new_policies"] == 1
    assert Decimal(rows[9]["claims_paid"]) == Decimal("6.0060")
    assert source == before


@pytest.fixture(scope="module")
def complete_source():
    catalog = json.loads((ROOT / "seminar_cases/bafin_2024_catalog.json").read_text(encoding="utf-8"))
    return build_case(catalog, "us_hyperscaler_outage", 5)


def test_contract_rejects_cycles_wrong_weights_unbound_source_and_prefix(complete_source):
    for edit in (lambda doc: doc["ict"]["assets"][1]["depends_on"].append("portal"),
                 lambda doc: doc["ict"]["resources"][0]["owners"][0].update(weight="0.9999"),
                 lambda doc: doc["reference_bundle"]["source_catalog"]["source_scope"].update(notice="changed"),
                 lambda doc: doc["events"][0].update(start_hour="96")):
        doc = deepcopy(complete_source)
        edit(doc)
        with pytest.raises(ContractError):
            validate(doc)


def test_api_fresh_digest_roundtrip_export_and_stale_rejection(complete_source):
    with TestClient(create_market_app()) as client:
        preset = client.post("/shock-case", json={"case_id": "us_hyperscaler_outage", "period_count": 5})
        assert preset.status_code == 200 and preset.json()["shock_bundle"] == complete_source
        calculated = client.post("/calculate", json=complete_source)
        assert calculated.status_code == 200, calculated.text[:200]
        body = calculated.json()
        assert body["transport"] == "ims.market-row-table.v1"
        aid = complete_source["model_input"]["insurers"][0]["insurer_id"]
        assert client.post("/export.xlsx", json={"source_input": complete_source, "insurer_id": aid}).status_code == 412
        headers = {"If-Match": calculated.headers["etag"]}
        restored = client.post("/source.json", json=complete_source, headers=headers)
        assert restored.status_code == 200 and restored.json() == complete_source
        exported = client.post("/export.xlsx", json={"source_input": complete_source, "insurer_id": aid}, headers=headers)
        assert exported.status_code == 200
        workbook = load_workbook(BytesIO(exported.content), read_only=True)
        assert {"ICT-Prozesse", "Kosten-einmal", "Lebens-Anträge", "Horizont-Rückstand"} <= set(workbook.sheetnames)


def test_capacity_and_fallback_are_order_independent():
    fallback = {"response_id": "Q", "kind": "fallback", "decision_period": 6, "lead_periods": 0, "duration_periods": 2,
                "asset_id": "primary", "replacement_asset_id": "independent", "value": "1"}
    staff = {**fallback, "response_id": "staff", "kind": "capacity", "value": "2"}
    for responses in ([fallback, staff], [staff, fallback]):
        worker = engine(process_bundle(events=[outage()], responses=responses))
        assert worker.capacity(worker.resources["shared"], Decimal(120))[0] == 2


def test_inactive_entry_capital_once_and_no_shock_keeps_precaution():
    catalog = json.loads((ROOT / "seminar_cases/bafin_2024_catalog.json").read_text(encoding="utf-8"))
    source = build_case(catalog, "google_motor_entry", 25)
    before = deepcopy(source)
    for side in ("baseline", "variant"):
        plan = ShockPlan(source, side)
        result = side_result(source["model_input"], side, {}, plan)
        entrant = [row for row in result["vu_rows"] if row["insurer_id"] == 1_000_000]
        assert all(not row["active_in_period"] and Decimal(row["premium_income"]) == 0 and Decimal(row["closing_assets"]) == 0 for row in entrant[:20])
        assert entrant[20]["active_in_period"]
        assert sum(Decimal(row["capital_contribution"]) for row in entrant) == 100_000
        assert entrant[20]["capital_contribution"] == "100000.0000"
        total = [row for row in result["market_rows"] if row["sector_id"] == "motor"]
        assert total[19]["active_vu_count"] + 1 == total[20]["active_vu_count"]
    assert source == before
    source["comparison"] = "no_shock_vs_shock"
    control = ShockPlan(source, "baseline")
    assert not control.active(1_000_000, 25)
    assert sum(Decimal(row["amount"]) for row in control.ledger.rows if row["cost_id"] == "decision_Preisantwort") == 100


def test_life_demand_is_one_pool_and_does_not_edit_old_policies():
    catalog = json.loads((ROOT / "seminar_cases/bafin_2024_catalog.json").read_text(encoding="utf-8"))
    source = build_case(catalog, "life_demand_shock", 25)
    before = deepcopy(source)
    plans = [ShockPlan(source, side) for side in ("baseline", "variant")]
    assert plans[0].life_rows[19]["willing"] == plans[1].life_rows[19]["willing"] == 4
    assert plans[0].life_rows[20]["willing"] == 1
    assert plans[1].life_rows[20]["willing"] == 2
    assert source == before
    source["comparison"] = "no_shock_vs_shock"
    control = ShockPlan(source, "baseline")
    assert control.life_rows[20]["willing"] == 4
    assert sum(Decimal(row["amount"]) for row in control.ledger.rows if row["cost_id"] == "decision_Vertriebsantwort") == 6


def test_contract_rejects_unfunded_response_and_hidden_reference_edit(complete_source):
    for edit in (
        lambda doc: doc["responses"]["variant"][0].update(asset_id="eu-keys"),
        lambda doc: doc["responses"]["variant"][0].update(insurer_ids=[doc["model_input"]["insurers"][0]["insurer_id"]]),
        lambda doc: doc["reference_bundle"]["model_input"]["insurers"][0].update(name="renamed"),
        lambda doc: doc["reference_bundle"]["model_input"]["damage_indicators"].__setitem__(0, "2"),
    ):
        source = deepcopy(complete_source)
        edit(source)
        with pytest.raises(ContractError):
            validate(source)


def test_budget_failure_is_atomic_and_returned_without_financial_rows(complete_source, monkeypatch):
    import ims.market.shock_runner as runner
    monkeypatch.setattr(runner, "MAX_RESULT_BYTES", 100)
    result = runner.calculate(complete_source)
    assert not result["valid"] and result["sides"] == {}
    assert result["content_digest"] is None
    assert not result["partial_result_returned"] and not result["writes_performed"]


def test_sequence_uses_half_open_edges_and_one_shared_budget():
    source = process_bundle(rate="2", events=[outage("120", "6"), {**outage("132", "6"), "event_id": "second"}])
    worker = engine(source)
    done = worker.period(6, [Job(f"A{i:02}", "U1", 1, "motor", "switch", 6) for i in range(40)], {1, 2})
    assert len(done) == 24  # 12 available hours × 2, not two separate queue budgets.
    assert worker.resource_rows[0]["capacity_work"] == "24.0000"
    assert [row["from_hour"] for row in worker.resource_rows[0]["capacity_segments"]] == ["120.0000", "126.0000", "132.0000", "138.0000"]
    assert worker.state("primary", Decimal(126))[0] == worker.state("primary", Decimal(138))[0] == 1
