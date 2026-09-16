from copy import deepcopy
import csv
from hashlib import sha256
from io import BytesIO, StringIO
import json
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

from ims.api.strategy_execution_period_chain_extended_probe import (
    parse_extended_probe_request,
    run_extended_probe,
)
from ims.api.strategy_execution_result_bundle import (
    BUNDLE_VERSION,
    COLUMNS,
    METRICS_VERSION,
    ResultBundleError,
    build_result_bundle,
)
from tests.test_strategy_execution_period_chain_extended_probe import _input


@pytest.fixture(scope="module")
def hundred_result(tmp_path_factory):
    root = tmp_path_factory.mktemp("result-bundle")
    db_path, payload = _input(root, 100)
    before = db_path.read_bytes()
    result = run_extended_probe(parse_extended_probe_request(payload), db_path=db_path)
    assert db_path.read_bytes() == before
    return result


def test_bundle_contains_identical_table_in_three_formats_with_full_evidence(hundred_result):
    bundle = build_result_bundle(hundred_result)
    with ZipFile(BytesIO(bundle)) as archive:
        assert set(archive.namelist()) == {
            "manifest.json", "ergebnis.json", "kennzahlen.csv", "kennzahlen.xlsx"
        }
        manifest = json.loads(archive.read("manifest.json"))
        payload = json.loads(archive.read("ergebnis.json"))
        assert manifest["schema_version"] == BUNDLE_VERSION
        assert payload["schema_version"] == METRICS_VERSION
        assert payload["source_result"] == hundred_result
        assert payload["columns"] == list(COLUMNS) == manifest["columns"]
        assert manifest["period_count"] == 100
        assert manifest["row_count"] == len(payload["rows"])
        assert manifest["row_count"] > 100
        assert manifest["effect_digest"] == hundred_result["effect_digest"]
        assert manifest["content_digest"] == hundred_result["period_chain_identity"]["content_digest"]
        assert manifest["result_persisted"] is False
        assert manifest["historical_full_equality_claim"] is False
        for name, entry in manifest["files"].items():
            data = archive.read(name)
            assert entry == {"sha256": "sha256:" + sha256(data).hexdigest(),
                             "byte_count": len(data)}

        csv_rows = list(csv.DictReader(StringIO(archive.read("kennzahlen.csv").decode("utf-8"))))
        assert csv_rows == payload["rows"]
        sheet = load_workbook(BytesIO(archive.read("kennzahlen.xlsx")), read_only=True).active
        values = sheet.iter_rows(values_only=True)
        assert next(values) == COLUMNS
        assert [dict(zip(COLUMNS, row)) for row in values] == payload["rows"]
        assert {int(row["period"]) for row in csv_rows} == set(range(1, 101))
        assert {row["stage"] for row in csv_rows} == {
            "applications", "state_before", "state_after"
        }
        assert any(row["sector_index"] == "0" for row in csv_rows)


@pytest.mark.parametrize("mutation", [
    lambda value: value.update(effect_digest="sha256:" + "0" * 64),
    lambda value: value["period_effects"].pop(),
    lambda value: value["period_effects"][4].update(period=6),
    lambda value: value["prefix_proof"].update(canonical_json_byte_equal=False),
    lambda value: value["period_effects"][0]["state_after"]["insurers"].append(
        {"insurer_id": 1, "premiums_current": "=1+1"}
    ),
])
def test_bundle_rejects_corrupt_or_incomplete_evidence(hundred_result, mutation):
    altered = deepcopy(hundred_result)
    mutation(altered)
    with pytest.raises(ResultBundleError):
        build_result_bundle(altered)


def test_bundle_rejects_nonfinite_and_formula_values_even_with_matching_digest(hundred_result):
    altered = deepcopy(hundred_result)
    altered["period_effects"][0]["state_after"]["insurers"][0]["premiums_current"] = "=1+1"
    worker_fields = set(altered) - {
        "effect_digest", "prefix_proof", "peak_worker_rss_bytes",
        "chain_payload_bytes", "result_payload_bytes", "wall_elapsed_seconds",
    }
    raw = json.dumps({key: altered[key] for key in worker_fields}, sort_keys=True,
                     ensure_ascii=True, allow_nan=False, separators=(",", ":")).encode("ascii")
    altered["effect_digest"] = "sha256:" + sha256(raw).hexdigest()
    altered["result_payload_bytes"] = len(raw)
    with pytest.raises(ResultBundleError) as caught:
        build_result_bundle(altered)
    assert caught.value.code == "metric_invalid"


def test_bundle_rejects_missing_provenance_and_oversized_output(hundred_result, monkeypatch):
    import ims.api.strategy_execution_result_bundle as module

    altered = deepcopy(hundred_result)
    altered["period_chain_identity"]["chain_id"] = "=cmd()"
    with pytest.raises(ResultBundleError):
        build_result_bundle(altered)

    monkeypatch.setattr(module, "_MAX_ROWS", 1)
    with pytest.raises(ResultBundleError) as caught:
        build_result_bundle(hundred_result)
    assert caught.value.code == "result_too_large"

    monkeypatch.setattr(module, "_MAX_ROWS", 50_000)
    monkeypatch.setattr(module, "_MAX_BUNDLE_BYTES", 1)
    with pytest.raises(ResultBundleError) as caught:
        build_result_bundle(hundred_result)
    assert caught.value.code == "result_too_large"
