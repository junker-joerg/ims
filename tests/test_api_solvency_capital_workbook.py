from io import BytesIO
import importlib

from openpyxl import load_workbook
import pytest
from starlette.testclient import TestClient

from ims.accounting.solvency_capital_readiness import build_solvency_capital_readiness
from ims.accounting.solvency_capital_workbook import build_solvency_capital_workbook
from ims.accounting.solvency_model_balance import build_solvency_model_balance
from ims.accounting.solvency_risk_aggregation import build_solvency_risk_aggregation
from ims.api.app import create_app
from tests.test_solvency_capital_readiness import _balance_input, _input


def test_workbook_contains_exact_model_values_and_blank_regulatory_cells() -> None:
    value = _input()
    readiness = build_solvency_capital_readiness(value).to_dict()
    balance = build_solvency_model_balance(_balance_input(value)).to_dict()
    aggregation = build_solvency_risk_aggregation(value["solvency_risk_aggregation_input"]).to_dict()
    contents = build_solvency_capital_workbook(value, readiness, balance, aggregation)
    workbook = load_workbook(BytesIO(contents), read_only=True, data_only=False)
    assert workbook.sheetnames == ["Modellwirkung", "Sparten", "Risiken", "Annahmen", "Regulatorik", "Herkunft"]
    model = list(workbook["Modellwirkung"].values)
    assert ("Modell", "Eigenmittel-Proxy", "319.0000", "Modellwaehrung, nicht EUR") in model
    assert ("Modell", "Netto-Stressverlust", "1.6000", "Kein SCR") in model
    assert ("Modell", "Verbleibender Eigenkapital-Proxy", "317.4000", "Keine Bilanzbuchung") in model
    assert ("Workshop", "Verlustgrenze eingehalten", "Ja", "Nicht regulatorisch") in model
    assert workbook["Modellwirkung"]["C2"].data_type == "s"
    sectors = list(workbook["Sparten"].values)
    assert ("Gesamt", "462.0000", "143.0000", "319.0000") in sectors
    assert [row[0] for row in sectors[1:5]] == ["Kfz", "Sach-Haftpflicht", "Leben", "Kranken"]
    risks = list(workbook["Risiken"].values)
    assert ("Marktwert Aktiva", "0.2500", "1.0000") in risks
    assert any(row[:2] == ("Summe vor Aggregation", "1.7750") for row in risks)
    assumptions = list(workbook["Annahmen"].values)
    assert ("Teilposition", "motor_assets", "base_amount", "1") in assumptions
    assert ("Modulsatz", "motor_assets", "stress_rate", "0.25") in assumptions
    assert ("Workshop-Grenze", "management", "max_net_model_stress_loss", "2") in assumptions
    blocked = list(workbook["Regulatorik"].values)
    assert [row[0] for row in blocked[1:6]] == [
        "Anrechenbare Eigenmittel", "SCR", "MCR", "SCR-Bedeckungsquote", "MCR-Bedeckungsquote",
    ]
    assert all(row[1] is None and "Nicht berechnet" in row[2] for row in blocked[1:6])
    origin = dict(list(workbook["Herkunft"].values)[1:])
    assert origin["Eingang-Digest"] == readiness["input_digest"]
    assert origin["Vier-Sparten-Bilanz-Digest"] == balance["source_content_digest"]
    assert origin["PR172-Modellbilanz-Digest"] == balance["content_digest"]
    assert origin["PR175-Stress-Digest"] == aggregation["content_digest"]
    assert origin["PR176-Kapitalansicht-Digest"] == readiness["content_digest"]
    workbook.close()


def test_workbook_rejects_invalid_or_mismatched_reports() -> None:
    value = _input()
    readiness = build_solvency_capital_readiness(value).to_dict()
    balance = build_solvency_model_balance(_balance_input(value)).to_dict()
    aggregation = build_solvency_risk_aggregation(value["solvency_risk_aggregation_input"]).to_dict()
    changed = dict(readiness, source_model_balance_content_digest="0" * 64)
    with pytest.raises(ValueError, match="Quellen-Digests"):
        build_solvency_capital_workbook(value, changed, balance, aggregation)
    changed = dict(readiness, regulatory_metrics=dict(readiness["regulatory_metrics"], scr="1.0000"))
    with pytest.raises(ValueError, match="gesperrt"):
        build_solvency_capital_workbook(value, changed, balance, aggregation)
    with pytest.raises(ValueError, match="gueltige"):
        build_solvency_capital_workbook(value, dict(readiness, valid=False), balance, aggregation)
    changed_input = _input()
    changed_input["management_limits"]["max_net_model_stress_loss"] = "3"
    with pytest.raises(ValueError, match="Eingang"):
        build_solvency_capital_workbook(changed_input, readiness, balance, aggregation)


@pytest.mark.parametrize("starlette_fallback", [False, True])
def test_workbook_api_is_read_only_and_bound_to_same_result(
    monkeypatch, tmp_path, starlette_fallback: bool,
) -> None:
    app_module = importlib.import_module("ims.api.app")
    if starlette_fallback:
        monkeypatch.setattr(app_module, "FastAPI", None)
    db_path = tmp_path / "unused.sqlite"
    monkeypatch.setenv("IMS_METADATA_DB", str(db_path))
    runner = importlib.import_module("ims.engine.explicit_period_runner")

    def reject_execution(*args: object, **kwargs: object) -> None:
        raise AssertionError("Kapital-Export darf keinen Runner starten")

    monkeypatch.setattr(runner, "run_loaded_explicit_period", reject_execution)
    client = TestClient(create_app(frontend_dist=tmp_path))
    value = _input()
    endpoint = "/api/accounting/solvency-capital-readiness"
    contract = client.get(endpoint + "-contract").json()
    assert contract["xlsx_endpoint"] == endpoint + ".xlsx"
    report = client.post(endpoint, json=value)
    export = client.post(contract["xlsx_endpoint"], json=value)
    assert report.status_code == export.status_code == 200
    assert export.headers["etag"] == report.headers["etag"]
    assert export.headers["cache-control"] == "no-store"
    assert export.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "ims-kapital-modell-vu-1-p-1.xlsx" in export.headers["content-disposition"]
    workbook = load_workbook(BytesIO(export.content), read_only=True)
    assert dict(list(workbook["Herkunft"].values)[1:])["PR176-Kapitalansicht-Digest"] == (
        report.json()["content_digest"]
    )
    workbook.close()
    assert db_path.exists() is False
    assert client.get(contract["xlsx_endpoint"]).status_code == 405

    value["management_limits"]["max_net_model_stress_loss"] = "bad"
    invalid = client.post(contract["xlsx_endpoint"], json=value)
    assert invalid.status_code == 422
    assert invalid.headers["cache-control"] == "no-store"
    assert invalid.json()["management_evaluation"] is None
    assert invalid.json()["regulatory_metrics"]["scr"] is None
    malformed = client.post(
        contract["xlsx_endpoint"], content="{", headers={"content-type": "application/json"},
    )
    assert malformed.status_code == 400
    assert malformed.json()["partial_result_returned"] is False
    assert db_path.exists() is False
