import csv
import importlib
from io import BytesIO, StringIO

import pytest
from openpyxl import load_workbook
from starlette.testclient import TestClient

from ims.api.app import create_app
from ims.api.management_case import MAX_INPUT_BYTES


@pytest.mark.parametrize("fallback", [False, True])
def test_common_hundred_api_and_three_same_digest_exports(monkeypatch, tmp_path, fallback):
    module = importlib.import_module("ims.api.app")
    if fallback: monkeypatch.setattr(module, "FastAPI", None)
    path = tmp_path / "not-created.sqlite"
    monkeypatch.setenv("IMS_METADATA_DB", str(path))
    client = TestClient(create_app(frontend_dist=tmp_path))
    base = "/api/management-case"
    source = client.post(base + "/workshop-case", json={"case_id": "inflation", "period_count": 100, "insurer_id": 1}).json()["source_input"]
    response = client.post(base + "/calculate", json=source)
    assert response.status_code == 200
    result, etag = response.json(), response.headers["etag"]
    assert etag == '"' + result["content_digest"] + '"'
    for extension in ("json", "csv", "xlsx"):
        exported = client.post(base + f"/export.{extension}", json=source, headers={"If-Match": etag})
        assert exported.status_code == 200 and exported.headers["etag"] == etag
        if extension == "json": assert exported.json() == result
        elif extension == "csv":
            rows = list(csv.DictReader(StringIO(exported.content.decode("utf-8-sig"))))
            assert len(rows) == 1000 and {row["content_digest"] for row in rows} == {result["content_digest"]}
            assert rows[499]["closing_equity"] == "87500.0000"
        else:
            book = load_workbook(BytesIO(exported.content), data_only=False)
            assert book["baseline"].max_row == 501
            assert book["Herkunft"]["B1"].value == result["content_digest"]
            assert book["baseline"]["O501"].value == "87500.0000"
    assert not path.exists()
    for route in ("workshop-case", "declare-sources", "calculate"):
        for content in (b"{", b" " * (MAX_INPUT_BYTES + 1)):
            invalid = client.post(base + "/" + route, content=content)
            assert invalid.status_code == 422
            assert invalid.json()["content_digest"] is None
            assert invalid.headers["cache-control"] == "no-store"
    assert client.post(base + "/export.json", json=source).status_code == 412
    source["sides"]["variant"]["four_sector_input"]["insurer_id"] = 2
    bad = client.post(base + "/calculate", json=source)
    assert bad.status_code == 422 and bad.json()["sides"] == {}
    assert not path.exists()
