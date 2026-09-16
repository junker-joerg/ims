import copy
import json
from decimal import Decimal, localcontext
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from ims.accounting.insurer_balance import (
    INSURER_BALANCE_INPUT_VERSION,
    INSURER_BALANCE_RESULT_VERSION,
    build_insurer_balance,
    insurer_balance_workbench_contract_payload,
)
from ims.accounting.insurer_balance_workbook import build_insurer_balance_workbook
from ims.accounting.model_balance_contract import BALANCE_FIELDS


FIXTURE = Path(__file__).parent / "fixtures" / "non_life_model_balance_v1.json"


def _input() -> dict:
    motor = json.loads(FIXTURE.read_text(encoding="utf-8"))
    property_liability = copy.deepcopy(motor)
    property_liability["sector_id"] = "property_liability"
    property_liability["opening"] = {
        "opening_cash": "50", "opening_claim_liability": "10", "opening_equity": "40",
    }
    property_liability["periods"][0].update({
        "premium_income": "10", "investment_income": "1", "claims_incurred": "3",
        "claims_paid": "2", "operating_expense": "2", "capital_contribution": "0",
        "capital_distribution": "0",
    })
    property_liability["periods"][1].update({
        "premium_income": "8", "investment_income": "0", "claims_incurred": "4",
        "claims_paid": "3", "operating_expense": "1", "capital_contribution": "0",
        "capital_distribution": "0",
    })
    return {
        "schema_version": INSURER_BALANCE_INPUT_VERSION,
        "insurer_id": 1,
        "sectors": [motor, property_liability],
    }


def test_consolidation_sums_each_field_and_preserves_balance_and_carryover() -> None:
    value = _input()
    before = copy.deepcopy(value)
    result = build_insurer_balance(value).to_dict()

    assert result["schema_version"] == INSURER_BALANCE_RESULT_VERSION
    assert result["valid"] is True
    assert result["insurer_id"] == 1
    assert result["period_count"] == 2
    assert len(result["content_digest"]) == 64
    assert result["source_kind"] == "explicit_scenario"
    assert result["historical_mapping_status"] == "unresolved"
    assert result["writes_performed"] is False
    assert result["runner_invoked"] is False
    assert result["simulation_performed"] is False
    assert result["historical_full_equality_claim"] is False
    assert value == before
    assert result == build_insurer_balance(value).to_dict()

    first, second = result["total_rows"]
    assert [item["sector_id"] for item in result["sectors"]] == ["motor", "property_liability"]
    assert (first["closing_cash"], first["closing_claim_liability"], first["closing_equity"]) == (
        "174.0000", "44.0000", "130.0000"
    )
    assert (second["closing_cash"], second["closing_claim_liability"], second["closing_equity"]) == (
        "177.0000", "41.0000", "136.0000"
    )
    assert second["opening_cash"] == first["closing_cash"]
    assert second["opening_claim_liability"] == first["closing_claim_liability"]
    assert second["opening_equity"] == first["closing_equity"]
    for index, total in enumerate(result["total_rows"]):
        motor = result["sectors"][0]["rows"][index]
        property_liability = result["sectors"][1]["rows"][index]
        for field in BALANCE_FIELDS:
            assert Decimal(total[field.field_id]) == (
                Decimal(motor[field.field_id]) + Decimal(property_liability[field.field_id])
            )
        for prefix in ("opening", "closing"):
            assert Decimal(total[f"{prefix}_cash"]) == (
                Decimal(total[f"{prefix}_claim_liability"]) + Decimal(total[f"{prefix}_equity"])
            )

    value["sectors"].reverse()
    assert build_insurer_balance(value).to_dict() == result


def test_hundred_periods_keep_two_period_prefix_and_decimal_precision() -> None:
    value = _input()
    prefix = build_insurer_balance(value).to_dict()["total_rows"]
    for sector in value["sectors"]:
        empty = {name: "0" for name in sector["periods"][0] if name != "period"}
        sector["periods"].extend({"period": period, **empty} for period in range(3, 101))
    with localcontext() as context:
        context.prec = 3
        result = build_insurer_balance(value).to_dict()
    assert result["valid"] is True
    assert result["period_count"] == 100
    assert result["total_rows"][:2] == prefix
    assert result["total_rows"][-1]["closing_cash"] == "177.0000"


def test_bad_sector_or_period_never_returns_partial_totals() -> None:
    scenarios = []
    mismatch = _input()
    mismatch["sectors"][1]["insurer_id"] = 2
    scenarios.append((mismatch, "insurer_mismatch"))
    duplicate = _input()
    duplicate["sectors"][1]["sector_id"] = "motor"
    scenarios.append((duplicate, "sector_pair_invalid"))
    short = _input()
    short["sectors"][1]["periods"].pop()
    scenarios.append((short, "period_count_mismatch"))
    late_error = _input()
    late_error["sectors"][1]["periods"][1]["claims_paid"] = "999"
    scenarios.append((late_error, "negative_closing_claim_liability"))
    bad_source = _input()
    bad_source["sectors"][0]["source_kind"] = "legacy"
    scenarios.append((bad_source, "contract_value_mismatch"))

    for value, code in scenarios:
        result = build_insurer_balance(value).to_dict()
        assert result["valid"] is False
        assert result["period_count"] == 0
        assert result["sectors"] == []
        assert result["total_rows"] == []
        assert result["content_digest"] is None
        assert code in {issue["code"] for issue in result["issues"]}


def test_unknown_top_field_and_unsupported_shape_are_rejected() -> None:
    value = _input()
    value["legacy_reserves"] = [100, 200]
    assert "field_unknown" in {
        issue["code"] for issue in build_insurer_balance(value).to_dict()["issues"]
    }
    assert build_insurer_balance(None).to_dict()["valid"] is False
    value = _input()
    value["sectors"] = value["sectors"][:1]
    assert "sector_count_invalid" in {
        issue["code"] for issue in build_insurer_balance(value).to_dict()["issues"]
    }


def test_xlsx_has_same_exact_values_and_source_boundary() -> None:
    report = build_insurer_balance(_input())
    payload = report.to_dict()
    workbook = load_workbook(BytesIO(build_insurer_balance_workbook(report)), read_only=True)
    try:
        assert workbook.sheetnames == ["Gesamt", "Kfz", "Sach-Haftpflicht", "Herkunft"]
        for sheet_name, rows in (
            ("Gesamt", payload["total_rows"]),
            ("Kfz", payload["sectors"][0]["rows"]),
            ("Sach-Haftpflicht", payload["sectors"][1]["rows"]),
        ):
            sheet = workbook[sheet_name]
            assert sheet["A2"].value == 1
            assert sheet["M2"].value == rows[0]["closing_cash"]
            assert sheet["M2"].data_type == "s"
            assert sheet["M3"].value == rows[1]["closing_cash"]
            assert sheet["M3"].data_type == "s"
        provenance = dict(workbook["Herkunft"].values)
        assert provenance["Ergebnis-Digest"] == payload["content_digest"]
        assert provenance["Historische Spartenbindung"] == "Offen"
        assert "Text" in provenance["Betragszellen"]
    finally:
        workbook.close()


def test_workbook_rejects_invalid_result_and_contract_states_bounds() -> None:
    value = _input()
    value["sectors"][0]["sector_id"] = "life"
    try:
        build_insurer_balance_workbook(build_insurer_balance(value))
    except ValueError:
        pass
    else:
        raise AssertionError("Fehlerhafte Bilanz darf keinen XLSX-Export erhalten")
    contract = insurer_balance_workbench_contract_payload()
    assert contract["sector_ids"] == ["motor", "property_liability"]
    assert contract["xlsx_amount_cells"] == "canonical_text"
    assert contract["xlsx_if_match_required"] is True
    assert contract["writes_enabled"] is False
    assert contract["runner_enabled"] is False
