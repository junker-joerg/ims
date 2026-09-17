import importlib
import json
import sqlite3
import time
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from starlette.testclient import TestClient

from ims.accounting.life_period_chain import run_life_policy_period_chain
from ims.api.life_result_delivery import (
    LIFE_RESULT_START_VERSION,
    digest,
    parse_start,
    persist_life_result,
)


FIXTURE = Path(__file__).parent / "fixtures" / "life_period_chain_v1.json"
BASE = "/api/accounting/life-period-chain"


def _input() -> dict:
    return json.loads(FIXTURE.read_text(encoding="utf-8"))


def _start(value: dict, preview: dict, *, key: str = "life-test-1") -> dict:
    return {
        "schema_version": LIFE_RESULT_START_VERSION,
        "life_chain_input": value,
        "expected_input_digest": preview["input_digest"],
        "expected_result_digest": preview["result_digest"],
        "idempotency_key": key,
        "explicit_storage_release": True,
    }


def _sheet_rows(sheet) -> list[dict]:
    values = list(sheet.values)
    return [dict(zip(values[0], row)) for row in values[1:]]


@pytest.mark.parametrize("fallback", [False, True])
def test_preview_store_replay_read_and_xlsx_share_exact_numbers(
    monkeypatch, tmp_path, fallback: bool,
) -> None:
    module = importlib.import_module("ims.api.app")
    if fallback:
        monkeypatch.setattr(module, "FastAPI", None)
    db_path = tmp_path / "metadata.sqlite"
    monkeypatch.setenv("IMS_METADATA_DB", str(db_path))
    client = TestClient(module.create_app(frontend_dist=tmp_path))
    value = _input()

    contract = client.get(f"{BASE}/contract")
    assert contract.status_code == 200
    assert contract.json()["timeout_seconds"] == 20.0
    preview = client.post(f"{BASE}/preview", json=value)
    assert preview.status_code == 200
    assert preview.json()["report"]["valid"] is True
    assert preview.json()["input_digest"] == digest(value)
    assert preview.json()["result_digest"] == digest(preview.json()["report"])
    assert db_path.exists() is False
    assert client.get(f"{BASE}/results").json()["results"] == []

    request = _start(value, preview.json())
    started = client.post(f"{BASE}/start", json=request)
    assert started.status_code == 200
    record = started.json()
    assert record["replayed"] is False
    assert record["writes_performed"] is True
    assert record["stored_result_verified"] is True
    assert record["result_digest"] == preview.json()["result_digest"]
    assert record["life_chain_input"] == value
    assert record["report"] == preview.json()["report"]
    assert db_path.is_file()
    history = client.get(f"{BASE}/results")
    assert history.status_code == 200
    assert history.json()["results"] == [{
        "result_id": record["result_id"],
        "insurer_id": 1,
        "period_count": 2,
        "stored_at": record["stored_at"],
        "input_digest": record["input_digest"],
        "result_digest": record["result_digest"],
        "record_digest": record["record_digest"],
    }]
    url = f"{BASE}/result/{record['result_id']}"
    fetched = client.get(url)
    assert fetched.status_code == 200
    assert fetched.json()["life_chain_input"] == value
    assert fetched.json()["report"] == record["report"]
    assert fetched.headers["etag"] == f'"{record["result_digest"]}"'

    def reject_recalculation(*args: object, **kwargs: object) -> None:
        raise AssertionError("Replay darf nicht neu rechnen")

    monkeypatch.setattr(module, "run_life_policy_period_chain", reject_recalculation)
    replayed = client.post(f"{BASE}/start", json=request)
    assert replayed.status_code == 200
    assert replayed.json()["replayed"] is True
    assert replayed.json()["writes_performed"] is False
    assert replayed.json()["result_id"] == record["result_id"]

    assert client.get(f"{url}.xlsx").status_code == 428
    assert client.get(f"{url}.xlsx", headers={"If-Match": '"sha256:bad"'}).status_code == 409
    exported = client.get(f"{url}.xlsx", headers={"If-Match": fetched.headers["etag"]})
    assert exported.status_code == 200
    assert exported.headers["etag"] == fetched.headers["etag"]
    assert "spreadsheetml.sheet" in exported.headers["content-type"]
    workbook = load_workbook(BytesIO(exported.content), read_only=True)
    try:
        periods = _sheet_rows(workbook["Perioden"])
        sources = _sheet_rows(workbook["Quellen"])
        policies = _sheet_rows(workbook["Policenbewegung"])
        assert len(periods) == 2
        for index, row in enumerate(periods):
            expected = record["report"]["rows"][index]
            for key, value in row.items():
                assert value == expected[key]
        assert sources[0]["investment_result"] == record["report"]["resolved_sources"][0]["investment_result"]
        assert policies[0]["renewal_premiums_collected"] == record["report"]["rows"][0]["policy_movements"][0]["renewal_premiums_collected"]
        assert workbook["Perioden"]["A2"].value == 1
        assert workbook["Perioden"]["B2"].data_type == "s"
        assert workbook["Herkunft"]["B2"].value == record["input_digest"]
        assert workbook["Herkunft"]["B3"].value == record["result_digest"]
    finally:
        workbook.close()
    assert client.get(f"{BASE}/preview").status_code == 405
    assert client.get(f"{BASE}/start").status_code == 405


@pytest.mark.parametrize("fallback", [False, True])
def test_start_failures_and_stored_tampering_are_atomic(monkeypatch, tmp_path, fallback: bool) -> None:
    module = importlib.import_module("ims.api.app")
    if fallback:
        monkeypatch.setattr(module, "FastAPI", None)
    db_path = tmp_path / "metadata.sqlite"
    monkeypatch.setenv("IMS_METADATA_DB", str(db_path))
    client = TestClient(module.create_app(frontend_dist=tmp_path))
    value = _input()
    preview = client.post(f"{BASE}/preview", json=value).json()
    request = _start(value, preview)

    bad = {**request, "expected_input_digest": "sha256:" + "0" * 64}
    assert client.post(f"{BASE}/start", json=bad).json()["code"] == "input_digest_mismatch"
    no_release = {**request, "explicit_storage_release": False}
    assert client.post(f"{BASE}/start", json=no_release).status_code == 403
    wrong_result = {**request, "expected_result_digest": "sha256:" + "0" * 64}
    mismatch = client.post(f"{BASE}/start", json=wrong_result)
    assert mismatch.status_code == 409
    assert mismatch.json()["code"] == "result_digest_mismatch"
    assert db_path.exists() is False

    invalid = _input()
    invalid["periods"][1]["policy_flows"].pop()
    rejected = client.post(f"{BASE}/preview", json=invalid)
    assert rejected.status_code == 422
    assert rejected.json()["rows"] == []
    assert rejected.json()["resolved_sources"] == []
    assert db_path.exists() is False
    assert client.post(f"{BASE}/preview", content="{", headers={"content-type": "application/json"}).status_code == 400

    stored = client.post(f"{BASE}/start", json=request).json()
    changed = _input()
    changed["periods"][0]["capital_contribution"] = "21"
    changed_preview = client.post(f"{BASE}/preview", json=changed).json()
    conflict = client.post(f"{BASE}/start", json=_start(changed, changed_preview))
    assert conflict.status_code == 409
    assert conflict.json()["code"] == "idempotency_key_conflict"

    with sqlite3.connect(db_path) as connection:
        connection.execute(
            "UPDATE life_chain_results SET stored_at = ? WHERE result_id = ?",
            ("2000-01-01T00:00:00+00:00", stored["result_id"]),
        )
    url = f"{BASE}/result/{stored['result_id']}"
    assert client.get(url).json()["code"] == "stored_life_result_corrupt"
    with sqlite3.connect(db_path) as connection:
        connection.execute(
            "UPDATE life_chain_results SET stored_at = ? WHERE result_id = ?",
            (stored["stored_at"], stored["result_id"]),
        )
        connection.execute(
            "UPDATE life_chain_results SET report_json = ? WHERE result_id = ?",
            ('{"valid":true,"rows":[]}', stored["result_id"]),
        )
    assert client.get(url).json()["code"] == "stored_life_result_corrupt"
    assert client.get(f"{BASE}/results").json()["code"] == "stored_life_result_corrupt"
    assert client.get(f"{url}.xlsx", headers={"If-Match": '"' + stored["result_digest"] + '"'}).status_code == 409
    assert client.post(f"{BASE}/start", json=request).status_code == 409


def test_start_requires_explicit_sqlite_and_timeout_never_writes(monkeypatch, tmp_path) -> None:
    module = importlib.import_module("ims.api.app")
    monkeypatch.delenv("IMS_METADATA_DB", raising=False)
    client = TestClient(module.create_app(frontend_dist=tmp_path))
    value = _input()
    preview = client.post(f"{BASE}/preview", json=value).json()
    request = _start(value, preview)
    unavailable = client.post(f"{BASE}/start", json=request)
    assert unavailable.status_code == 503
    assert unavailable.json()["code"] == "explicit_sqlite_storage_required"

    db_path = tmp_path / "metadata.sqlite"
    monkeypatch.setenv("IMS_METADATA_DB", str(db_path))
    monkeypatch.setattr(module, "LIFE_API_TIMEOUT_SECONDS", 0.01)
    original = module.run_life_policy_period_chain

    def wait_for_cancel(value, *, should_cancel):
        while not should_cancel():
            time.sleep(0.001)
        return original(value, should_cancel=should_cancel)

    monkeypatch.setattr(module, "run_life_policy_period_chain", wait_for_cancel)
    client = TestClient(module.create_app(frontend_dist=tmp_path))
    timeout = client.post(f"{BASE}/start", json=request)
    assert timeout.status_code == 504
    assert timeout.json()["code"] == "life_calculation_timeout"
    assert db_path.exists() is False


def test_parallel_identical_starts_create_one_verified_record(tmp_path) -> None:
    value = _input()
    report = run_life_policy_period_chain(value).to_dict()
    request = parse_start(_start(value, {
        "input_digest": digest(value), "result_digest": digest(report),
    }))
    db_path = tmp_path / "metadata.sqlite"
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(
            lambda _: persist_life_result(db_path, request, report), range(2),
        ))
    assert {result["result_id"] for result in results} == {results[0]["result_id"]}
    assert sorted(result["writes_performed"] for result in results) == [False, True]
    with sqlite3.connect(db_path) as connection:
        assert connection.execute("SELECT COUNT(*) FROM life_chain_results").fetchone()[0] == 1


def test_small_100_period_life_case_is_stored_and_exportable(monkeypatch, tmp_path) -> None:
    module = importlib.import_module("ims.api.app")
    monkeypatch.setenv("IMS_METADATA_DB", str(tmp_path / "metadata.sqlite"))
    client = TestClient(module.create_app(frontend_dist=tmp_path))
    value = _input()
    value["opening"] = {
        "opening_active_policies": 1, "opening_backing_assets": "120",
        "opening_guarantee_liability": "100", "opening_equity": "20",
        "cohorts": [{
            "cohort_id": "A", "issue_period": 0, "issue_term_periods": 100,
            "remaining_periods": 100, "active_policies": 1,
            "guaranteed_rate_per_period": "0", "guarantee_liability": "100",
        }],
        "policies": [{
            "policy_id": "A1", "cohort_id": "A", "issue_term_periods": 100,
            "remaining_periods": 100, "guaranteed_rate_per_period": "0",
            "guarantee_liability": "100",
        }],
    }
    value["assumptions"]["period_count"] = 100
    value["assumptions"]["investment"] = {
        "mode": "insurer_rule_on_opening_backing_assets",
        "windows": [{"start_period": 1, "end_period": 100, "rate_per_period": "0"}],
    }
    value["assumptions"]["mortality"] = {
        "mode": "explicit_death_policy_ids",
        "periods": [{"period": period, "death_policy_ids": []} for period in range(1, 101)],
    }
    value["periods"] = [{
        "period": period,
        "policy_flows": [{
            "policy_id": "A1", "renewal_premiums_collected": "0",
            "renewal_liability_allocation": "0", "death_benefit_if_death": "0",
            "maturity_benefit_if_due": "100" if period == 100 else "0",
        }],
        "new_business": [], "operating_expense_paid": "0",
        "capital_contribution": "0", "capital_distribution": "0",
    } for period in range(1, 101)]
    preview = client.post(f"{BASE}/preview", json=value)
    assert preview.status_code == 200
    started = client.post(f"{BASE}/start", json=_start(value, preview.json(), key="life-100-1"))
    assert started.status_code == 200
    record = started.json()
    assert record["report"]["calculated_period_count"] == 100
    assert record["report"]["rows"][-1]["closing_backing_assets"] == "20.0000"
    result_url = f"{BASE}/result/{record['result_id']}"
    assert client.get(result_url).json()["report"] == record["report"]
    download = client.get(f"{result_url}.xlsx", headers={"If-Match": f'"{record["result_digest"]}"'})
    assert download.status_code == 200
    workbook = load_workbook(BytesIO(download.content), read_only=True)
    try:
        assert len(_sheet_rows(workbook["Perioden"])) == 100
        assert _sheet_rows(workbook["Perioden"])[-1]["closing_backing_assets"] == "20.0000"
    finally:
        workbook.close()
