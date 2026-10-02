"""AP6 accepted reference scope: conserved facts, explicit assumptions, fresh exports."""
from copy import deepcopy
from decimal import Decimal
from io import BytesIO
import json
from pathlib import Path
import unittest

from openpyxl import load_workbook
from starlette.testclient import TestClient

from ims.api.market import create_market_app
from ims.market.reference import CATALOG_DIGEST, build_bundle, calculate, reference_values, settings_checked
from ims.market.runner import digest, FINANCIAL_FIELDS
from ims.market.transport import wire_payload
from ims.strategies.modern_bridge import ContractError

ROOT = Path(__file__).resolve().parents[1]


class ReferenceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.catalog = json.loads((ROOT / "seminar_cases/bafin_2024_catalog.json").read_text(encoding="utf-8"))

    def checked(self, bundle):
        result = calculate(bundle)
        self.assertTrue(result["valid"], result["issues"])
        return result

    def test_catalog_matches_pinned_audit_and_ids_are_not_ranks(self):
        self.assertEqual(digest(self.catalog), CATALOG_DIGEST)
        self.assertEqual(len({g["insurer_id"] for g in self.catalog["groups"]}), 145)
        self.assertEqual(len({g["group_id"] for g in self.catalog["groups"]}), 145)
        self.assertFalse(self.catalog["source_scope"]["german_direct_market_complete"])
        self.assertEqual(self.catalog["source_scope"]["nonlife_product_breakdown"], "unknown")
        bundle = build_bundle(self.catalog, 5)
        self.assertEqual(len(bundle["model_input"]["insurers"]), 40)
        self.assertNotEqual(bundle["model_input"]["insurers"][0]["insurer_id"], 1)

    def test_disjoint_rest_and_small_mapping_hand_case(self):
        bundle = build_bundle(self.catalog, 5)
        ref = reference_values(self.catalog, bundle["overrides"], bundle["workshop"])
        self.assertEqual(Decimal(ref["universe_total_million_eur"]), Decimal("267483.206"))
        self.assertEqual(Decimal(ref["selected_total_million_eur"]), Decimal("254270.039"))
        self.assertEqual(Decimal(ref["rest"]["unselected_groups_million_eur"]), Decimal("13213.167"))
        self.assertEqual(Decimal(ref["rest"]["selected_unmodeled_million_eur"]), Decimal("23433.6"))
        self.assertEqual(Decimal(ref["modeled_selected_million_eur"]) + sum(Decimal(v) for k, v in ref["rest"].items() if k.endswith("million_eur")), Decimal(ref["universe_total_million_eur"]))
        allianz = next(g for g in ref["groups"] if g["name"] == "Allianz")
        self.assertEqual(Decimal(allianz["model_basis_million_eur"]["motor"]), Decimal("7966"))
        self.assertEqual(Decimal(allianz["unmodeled_million_eur"]), Decimal("3983"))
        actor = next(a for a in bundle["model_input"]["insurers"] if a["insurer_id"] == allianz["insurer_id"])
        self.assertEqual(actor["sectors"]["motor"]["opening"], {"assets": "7966.0000", "liabilities": "0", "equity": "7966.0000"})
        cohort = next(c for c in bundle["model_input"]["customer_groups"] if c["initial_insurer_id"] == actor["insurer_id"] and c["sector_id"] == "motor")
        self.assertEqual(cohort["quantity"], "26.5533")
        self.assertEqual(cohort["risk_periods"][0]["loss"], "47.7959")
        result = self.checked(bundle)
        record = next(r for r in result["sides"]["baseline"]["vu_rows"] if r["insurer_id"] == actor["insurer_id"] and r["sector_id"] == "motor" and r["period"] == 1)
        self.assertEqual(record["premium_income"], "79.6599")
        self.assertEqual(record["closing_assets"], "7997.8640")

    def test_absent_source_sector_is_not_filled_with_synthetic_business(self):
        bundle = build_bundle(self.catalog, 5)
        talanx = next(a for a in bundle["model_input"]["insurers"] if a["name"].startswith("Talanx"))
        self.assertNotIn("health", talanx["sectors"])
        self.assertIn("life", talanx["sectors"])
        self.assertTrue(any(r["raw_value"] == "-" and r["premium_million_eur"] == "0" for r in self.catalog["entities"]))

    def test_override_can_change_selection_but_not_original_or_stable_ids(self):
        before = deepcopy(self.catalog)
        itzehoer = next(r for r in self.catalog["entities"] if r["group_assertion"] == "Itzehoer" and r["source_sector"] == "Schaden/Unfall")
        bundle = build_bundle(self.catalog, 5, overrides=[{"entity_row": itzehoer["workbook_row"], "premium_million_eur": "1000", "note": "Handfall, ausdrückliche Workshop-Annahme"}])
        ref = reference_values(self.catalog, bundle["overrides"], bundle["workshop"])
        self.assertIn("Itzehoer", {g["name"] for g in ref["groups"]})
        self.assertNotIn("Münchener Verein", {g["name"] for g in ref["groups"]})
        actor = next(a for a in bundle["model_input"]["insurers"] if a["name"] == "Itzehoer")
        registered = next(g for g in before["groups"] if g["name"] == "Itzehoer")
        self.assertEqual(actor["insurer_id"], registered["insurer_id"])
        self.assertEqual(bundle["source_catalog"], before)
        self.assertEqual(self.catalog, before)
        self.assertFalse(ref["german_direct_selection_verified"])

    def test_mix_rounding_does_not_destroy_exact_reference_conservation(self):
        bundle = build_bundle(self.catalog, 5)
        bundle["workshop"]["nonlife_mix"] = {"motor": "0.3333", "property_liability": "0.3333", "unmodeled": "0.3334"}
        ref = reference_values(self.catalog, [{"entity_row": 5, "premium_million_eur": "23896.0001", "note": "Präzisionshandfall"}], bundle["workshop"])
        sums = sum(sum(Decimal(v) for v in g["model_basis_million_eur"].values()) + Decimal(g["unmodeled_million_eur"]) for g in ref["groups"])
        self.assertEqual(sums, Decimal(ref["selected_total_million_eur"]))

    def test_invalid_mix_and_unexplained_or_duplicate_overrides_fail(self):
        settings = build_bundle(self.catalog, 5)["workshop"]
        settings["nonlife_mix"]["motor"] = "0.5"
        with self.assertRaisesRegex(ContractError, "zusammen 1"):
            settings_checked(settings)
        for edits in ([{"entity_row": 5, "premium_million_eur": "1", "note": ""}],
                      [{"entity_row": 5, "premium_million_eur": "NaN", "note": "Test"}],
                      [{"entity_row": 5, "premium_million_eur": "1", "note": "Test"}] * 2):
            with self.assertRaises(ContractError):
                build_bundle(self.catalog, 5, overrides=edits)

    def test_tampered_original_and_wrong_id_bindings_return_no_partial_result(self):
        bundle = build_bundle(self.catalog, 5)
        bundle["source_catalog"]["entities"][0]["premium_million_eur"] = "100"
        result = calculate(bundle)
        self.assertFalse(result["valid"])
        self.assertEqual(result["sides"], {})
        self.assertIn("Originaler Faktenkatalog", result["issues"][0]["message"])
        bundle = build_bundle(self.catalog, 5)
        bundle["model_input"]["insurers"][0]["insurance_group_id"] = "false_group"
        self.assertFalse(calculate(bundle)["valid"])

    def test_strategy_edits_are_declared_as_custom_model_and_digest_changes(self):
        bundle = build_bundle(self.catalog, 5)
        original = self.checked(bundle)
        bundle["model_input"]["families"][0]["parameters"]["price"] = "3.2"
        edited = self.checked(bundle)
        self.assertNotEqual(original["content_digest"], edited["content_digest"])
        self.assertEqual(edited["reference"]["model_binding"], "custom_workshop_model")
        self.assertEqual(original["reference"]["groups"], edited["reference"]["groups"])

    def test_api_portable_bundle_and_single_vu_excel_keep_same_sources(self):
        with TestClient(create_market_app()) as client:
            reply = client.post("/reference-case", json={"period_count": 5, "workshop": None, "overrides": None})
            self.assertEqual(reply.status_code, 200)
            bundle = reply.json()["source_bundle"]
            reply = client.post("/calculate", json=bundle)
            self.assertEqual(reply.status_code, 200)
            result, etag = reply.json(), reply.headers["etag"]
            aid = bundle["model_input"]["insurers"][0]["insurer_id"]
            source = client.post("/source.json", json=bundle, headers={"If-Match": etag})
            self.assertEqual(source.json(), bundle)
            report = client.post("/export.xlsx", json={"source_input": bundle, "insurer_id": aid}, headers={"If-Match": etag})
            self.assertEqual(report.status_code, 200)
            book = load_workbook(BytesIO(report.content), data_only=True)
            self.assertIn("BaFin-Quellenwerte", book.sheetnames)
            self.assertIn("Workshop-Annahmen", book.sheetnames)
            chunks = [r[1] for r in book["Herkunft"].values if r[0] == "source_bundle_json_chunk"]
            self.assertEqual(json.loads("".join(chunks)), bundle)
            self.assertTrue(all(c.data_type != "f" for s in book for r in s for c in r))
            bundle["workshop"]["note"] += " Änderung"
            stale = client.post("/export.xlsx", json={"source_input": bundle, "insurer_id": aid}, headers={"If-Match": etag})
            self.assertEqual(stale.status_code, 412)

    def test_40_group_100_period_run_conserves_money_and_prefix(self):
        short = self.checked(build_bundle(self.catalog, 5))
        long = self.checked(build_bundle(self.catalog, 100))
        for side in ("baseline", "variant"):
            self.assertEqual(short["sides"][side]["vu_rows"], [r for r in long["sides"][side]["vu_rows"] if r["period"] <= 5])
            for p in (1, 6, 100):
                total = next(r for r in long["sides"][side]["market_rows"] if r["period"] == p and r["sector_id"] == "total")
                actors = [r for r in long["sides"][side]["vu_rows"] if r["period"] == p]
                for field in FINANCIAL_FIELDS:
                    self.assertEqual(Decimal(total[field]), sum(Decimal(r[field]) for r in actors))
                self.assertEqual(Decimal(total["closing_assets"]), Decimal(total["closing_liabilities"]) + Decimal(total["closing_equity"]))
        self.assertEqual(long["source_bundle"]["source_catalog"], self.catalog)
        self.assertLess(len(json.dumps(wire_payload(long)).encode()), 64 * 1024 * 1024)


if __name__ == "__main__":
    unittest.main()
