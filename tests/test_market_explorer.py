"""Independent arithmetic and fresh boundary checks for the read-only AP8 layer."""
from copy import deepcopy
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
import threading
import unittest

from openpyxl import load_workbook
from starlette.testclient import TestClient

from ims.api.market import create_market_app
from ims.market.explorer import effective_switches, family_statistics, project_result, share_statistics
from ims.market.presets import workshop_case
from ims.market.runner import FINANCIAL_FIELDS, calculate
from ims.market.transport import unpack_rows, wire_payload
from ims.strategies.modern_bridge import ContractError


def financial(aid, assets, profit, premium):
    return {**{key: "0.0000" for key in FINANCIAL_FIELDS}, "insurer_id": aid,
            "opening_assets": assets, "period_profit": profit, "premium_income": premium}


class ExplorerTests(unittest.TestCase):
    def test_independent_90_10_weight_and_hhi_handcase(self):
        rows = [financial(1, "90", "9", "90"), financial(2, "10", "3", "10")]
        stats = family_statistics(rows)
        self.assertEqual(stats["weighted_profit_percent"], "12.0000")
        self.assertEqual((stats["individual_min_percent"], stats["individual_max_percent"]), ("10.0000", "30.0000"))
        shares = share_statistics(rows)
        self.assertEqual(shares["hhi"], "8200.0000")
        self.assertEqual([r["share_percent"] for r in shares["shares"]], ["90.0000", "10.0000"])

    def test_zero_negative_and_half_even_ratios(self):
        self.assertIsNone(share_statistics([financial(1, "0", "0", "0")])["hhi"])
        self.assertIsNone(share_statistics([financial(1, "1", "0", "-1"), financial(2, "1", "0", "2")])["hhi"])
        stats = family_statistics([financial(1, "0", "0", "0")])
        self.assertIsNone(stats["weighted_profit_percent"])
        self.assertEqual(stats["valid_individual_count"], 0)
        self.assertEqual(family_statistics([financial(1, "2000000", "1", "1")])["weighted_profit_percent"], "0.0000")
        self.assertEqual(family_statistics([financial(1, "2000000", "3", "3")])["weighted_profit_percent"], "0.0002")

    def test_multi_sector_insurer_is_summed_before_concentration_and_activity(self):
        rows = [financial(1, "1", "0", "30"), financial(1, "1", "0", "60"), financial(2, "0", "0", "10")]
        rows[-1]["active_in_period"] = False
        stats = share_statistics(rows)
        self.assertEqual((stats["registered_vu_count"], stats["active_vu_count"], stats["hhi"]), (2, 1, "8200.0000"))

    def test_actual_switches_exclude_opening_and_pending_and_keep_uninsured(self):
        rows = [{"period": p, "sector_id": s, "previous_insurer_id": a, "insurer_id": b, "quantity": q}
                for p, s, a, b, q in [(1,"motor",None,1,"10"),(6,"motor",1,2,"10"),
                                    (7,"motor",2,None,"10"),(8,"property_liability",None,1,"2"),(9,"motor",1,1,"10")]]
        selected = effective_switches(rows)
        self.assertEqual([r["period"] for r in selected], [6,7,8])
        self.assertEqual([r["sector_id"] for r in selected], ["motor","motor","property_liability"])
        selected[0]["quantity"] = "999"
        self.assertEqual(rows[1]["quantity"], "10")

    def test_pure_projection_preserves_core_rng_financials_and_dynamic_families(self):
        source = workshop_case("market", 10, 3)
        original = deepcopy(source)
        result = calculate(source)
        snapshot = deepcopy(result)
        view = project_result(result)
        self.assertEqual(source, original)
        self.assertEqual(result, snapshot)
        self.assertEqual(view["model_result_digest"], result["content_digest"])
        self.assertNotEqual(view["content_digest"], result["content_digest"])
        for side in ("baseline", "variant"):
            for projected, actual in zip(view["sides"][side]["financial_rows"], result["sides"][side]["vu_rows"]):
                self.assertEqual({k: projected[k] for k in FINANCIAL_FIELDS}, {k: actual[k] for k in FINANCIAL_FIELDS})
                self.assertNotIn("draws", projected)
            for family in view["sides"][side]["family_rows"]:
                actual = [r for r in result["sides"][side]["vu_rows"] if r["period"] == family["period"]
                          and r["family_id"] == family["family_id"] and (family["sector_id"] == "total" or r["sector_id"] == family["sector_id"])]
                self.assertEqual(family["member_ids"], sorted({r["insurer_id"] for r in actual}))
        packed = wire_payload(view)
        self.assertEqual({s:{k:unpack_rows(v) for k,v in ts.items()} for s,ts in packed["sides"].items()}, view["sides"])
        view["actors"][0]["name"] = "Changed"
        self.assertEqual(result, snapshot)

    def test_once_cost_book_bridge_and_broken_book_atomic(self):
        result = calculate(workshop_case("market", 10, 3))
        view = project_result(result)
        costs = [r for r in view["sides"]["variant"]["financial_rows"] if Decimal(r["measure_cost"]) > 0]
        self.assertTrue(costs)
        for row in costs:
            self.assertEqual(Decimal(row["closing_equity"]), Decimal(row["opening_equity"])+Decimal(row["period_profit"])+Decimal(row["capital_contribution"])-Decimal(row["capital_distribution"]))
        result["sides"]["variant"]["vu_rows"][0]["period_profit"] = "1.0000"
        with self.assertRaisesRegex(ContractError, "Buchungsidentität"):
            project_result(result)
        with patch("ims.market.explorer.MAX_RESULT_BYTES", 1):
            with self.assertRaisesRegex(ContractError, "40 MiB"):
                project_result(calculate(workshop_case("switch", 10, 2)))

    def test_fresh_view_and_core_export_digest_are_distinct_and_reimport_exact(self):
        source = workshop_case("switch", 10, 2)
        with TestClient(create_market_app()) as client:
            reply = client.post("/explore", json=source)
            self.assertEqual(reply.status_code, 200)
            view = reply.json()
            self.assertEqual(reply.headers["etag"], '"'+view["content_digest"]+'"')
            core_etag = '"'+view["model_result_digest"]+'"'
            self.assertEqual(client.post("/calculate", json=source).headers["etag"], core_etag)
            self.assertEqual(client.post("/source.json", json=source, headers={"If-Match":reply.headers["etag"]}).status_code, 412)
            portable = client.post("/source.json", json=source, headers={"If-Match":core_etag})
            self.assertEqual(portable.json(), source)
            self.assertEqual(client.post("/explore", json=portable.json()).headers["etag"], reply.headers["etag"])
            bookreply = client.post("/export.xlsx", json={"source_input":source,"insurer_id":2}, headers={"If-Match":core_etag})
            book = load_workbook(BytesIO(bookreply.content), read_only=True)
            records = list(book["Variant"].values) if "Variant" in book.sheetnames else list(book["Variante"].values)
            self.assertEqual({dict(zip(records[0], r))["insurer_id"] for r in records[1:]}, {"2"})
            source["seed"] += 1
            self.assertEqual(client.post("/source.json", json=source, headers={"If-Match":core_etag}).status_code, 412)

    def test_explore_shares_existing_gate_and_error_is_atomic(self):
        started, release = threading.Event(), threading.Event()
        source = workshop_case("switch", 10, 2)
        original = calculate
        def blocking(value):
            started.set(); release.wait(10); return original(value)
        app = create_market_app()
        with TestClient(app) as client, patch("ims.api.market.calculate", blocking):
            responses = []
            thread = threading.Thread(target=lambda: responses.append(client.post("/explore", json=source)))
            thread.start()
            try:
                self.assertTrue(started.wait(5))
                self.assertEqual(client.post("/calculate", json=source).status_code, 409)
            finally:
                release.set(); thread.join(10)
            self.assertEqual(responses[0].status_code, 200)
            self.assertEqual(client.post("/explore", json={}).status_code, 422)
        with TestClient(app) as client, patch("ims.market.explorer.MAX_RESULT_BYTES", 1):
            bad = client.post("/explore", json=source)
            self.assertEqual(bad.status_code, 422)
            self.assertFalse(bad.json()["valid"])
            self.assertFalse(bad.json().get("sides"))


if __name__ == "__main__":
    unittest.main()
