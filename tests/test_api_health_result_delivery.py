import csv
import importlib
import json
import sqlite3
import time
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO, StringIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from starlette.testclient import TestClient

from ims.accounting.health_period_chain import run_health_period_chain
from ims.api.health_result_delivery import (
    HEALTH_RESULT_START_VERSION,
    digest,
    parse_start,
    persist_health_result,
)


FIXTURE = Path(__file__).parent / "fixtures" / "health_period_chain_v1.json"
BASE = "/api/accounting/health-period-chain"


def _input() -> dict:
    return json.loads(FIXTURE.read_text(encoding="utf-8"))


def _start(value: dict, preview: dict, *, key: str = "health-test-1") -> dict:
    return {
        "schema_version": HEALTH_RESULT_START_VERSION,
        "health_chain_input": value,
        "expected_input_digest": preview["input_digest"],
        "expected_result_digest": preview["result_digest"],
        "idempotency_key": key,
        "explicit_storage_release": True,
    }


def _horizon_100() -> dict:
    value = _input()
    value["sources"]["period_count"] = 100
    for field in ("new_business", "exits"):
        value["sources"][field]["periods"] = [
            {"period": period, "count": 0} for period in range(1, 101)
        ]
    for field, amount in (("pricing", "2.00"), ("benefits", "1.00")):
        value["sources"][field]["windows"] = [
            {"start_period": 1, "end_period": 100, "amount_per_opening_policy": amount}
        ]
    value["periods"] = [{
        "period": period,
        "benefits_paid": "10.00",
        "investment_result": "0.00",
        "operating_expense_paid": "0.00",
        "capital_contribution": "0.00",
        "capital_distribution": "0.00",
    } for period in range(1, 101)]
    return value


def _sheet_rows(sheet) -> list[dict]:
    values = list(sheet.values)
    return [dict(zip(values[0], row)) for row in values[1:]]


@pytest.mark.parametrize("fallback", [False, True])
def test_preview_store_replay_read_and_three_exports_share_provenance(
    monkeypatch, tmp_path, fallback: bool,
) -> None:
    module = importlib.import_module("ims.api.app")
    if fallback:
        monkeypatch.setattr(module, "FastAPI", None)
    db_path = tmp_path / "metadata.sqlite"
    monkeypatch.setenv("IMS_METADATA_DB", str(db_path))
    client = TestClient(module.create_app(frontend_dist=tmp_path))
    value = _input()

    contract = client.get(f"{BASE}/contract")
    assert contract.status_code == 200
    assert contract.json()["timeout_seconds"] == 20.0
    assert contract.json()["export_if_match_required"] is True
    preview = client.post(f"{BASE}/preview", json=value)
    assert preview.status_code == 200
    assert preview.json()["report"]["calculated_period_count"] == 2
    assert preview.json()["input_digest"] == digest(value)
    assert preview.json()["result_digest"] == digest(preview.json()["report"])
    assert preview.json()["writes_performed"] is False
    assert db_path.exists() is False
    assert client.get(f"{BASE}/results").json()["results"] == []

    request = _start(value, preview.json())
    started = client.post(f"{BASE}/start", json=request)
    assert started.status_code == 200
    record = started.json()
    assert record["replayed"] is False
    assert record["writes_performed"] is True
    assert record["stored_result_verified"] is True
    assert record["report"] == preview.json()["report"]
    assert record["health_chain_input"] == value
    assert db_path.is_file()
    history = client.get(f"{BASE}/results")
    assert history.status_code == 200
    assert history.json()["results"] == [{
        "result_id": record["result_id"],
        "insurer_id": 1,
        "scenario_id": "seminar_health_01",
        "variant_id": "baseline",
        "period_count": 2,
        "stored_at": record["stored_at"],
        "input_digest": record["input_digest"],
        "result_digest": record["result_digest"],
        "record_digest": record["record_digest"],
    }]
    url = f"{BASE}/result/{record['result_id']}"
    fetched = client.get(url)
    assert fetched.status_code == 200
    assert fetched.json()["report"] == record["report"]
    assert fetched.headers["etag"] == f'"{record["result_digest"]}"'

    def reject_recalculation(*args: object, **kwargs: object) -> None:
        raise AssertionError("Replay darf nicht neu rechnen")

    monkeypatch.setattr(module, "run_health_period_chain", reject_recalculation)
    replayed = client.post(f"{BASE}/start", json=request)
    assert replayed.status_code == 200
    assert replayed.json()["replayed"] is True
    assert replayed.json()["writes_performed"] is False
    assert replayed.json()["result_id"] == record["result_id"]

    for extension in ("csv", "json", "xlsx"):
        assert client.get(f"{url}.{extension}").status_code == 428
        assert client.get(
            f"{url}.{extension}", headers={"If-Match": '"sha256:bad"'},
        ).status_code == 409
    headers = {"If-Match": fetched.headers["etag"]}
    exported_json = client.get(f"{url}.json", headers=headers)
    assert exported_json.status_code == 200
    assert exported_json.json() == fetched.json()
    assert exported_json.headers["etag"] == headers["If-Match"]
    assert "attachment" in exported_json.headers["content-disposition"]

    exported_csv = client.get(f"{url}.csv", headers=headers)
    assert exported_csv.status_code == 200
    csv_rows = list(csv.DictReader(StringIO(exported_csv.content.decode("utf-8"))))
    assert len(csv_rows) == 2
    for index, row in enumerate(csv_rows):
        for field in ("result_id", "input_digest", "result_digest", "record_digest"):
            assert row[field] == record[field]
        assert row["scenario_id"] == "seminar_health_01"
        assert row["variant_id"] == "baseline"
        for field, expected in record["report"]["rows"][index].items():
            assert row[field] == str(expected)

    exported_xlsx = client.get(f"{url}.xlsx", headers=headers)
    assert exported_xlsx.status_code == 200
    workbook = load_workbook(BytesIO(exported_xlsx.content), read_only=True)
    try:
        periods = _sheet_rows(workbook["Perioden"])
        assert len(periods) == 2
        for index, row in enumerate(periods):
            for field, expected in record["report"]["rows"][index].items():
                assert row[field] == expected
        assert workbook["Perioden"]["B2"].data_type == "s"
        provenance = dict(workbook["Herkunft"].values)
        assert provenance["Eingabe-Digest"] == record["input_digest"]
        assert provenance["Ergebnis-Digest"] == record["result_digest"]
        assert provenance["Datensatz-Digest"] == record["record_digest"]
        assert provenance["Szenario"] == "seminar_health_01"
        assert provenance["Variante"] == "baseline"
    finally:
        workbook.close()
    assert client.get(f"{BASE}/preview").status_code == 405
    assert client.get(f"{BASE}/start").status_code == 405


@pytest.mark.parametrize("fallback", [False, True])
def test_invalid_release_digest_and_stored_tampering_are_atomic(
    monkeypatch, tmp_path, fallback: bool,
) -> None:
    module = importlib.import_module("ims.api.app")
    if fallback:
        monkeypatch.setattr(module, "FastAPI", None)
    db_path = tmp_path / "metadata.sqlite"
    monkeypatch.setenv("IMS_METADATA_DB", str(db_path))
    client = TestClient(module.create_app(frontend_dist=tmp_path))
    value = _input()
    preview = client.post(f"{BASE}/preview", json=value).json()
    request = _start(value, preview)

    wrong_input = {**request, "expected_input_digest": "sha256:" + "0" * 64}
    assert client.post(f"{BASE}/start", json=wrong_input).json()["code"] == "input_digest_mismatch"
    assert client.post(f"{BASE}/start", json={**request, "explicit_storage_release": False}).status_code == 403
    wrong_result = {**request, "expected_result_digest": "sha256:" + "0" * 64}
    assert client.post(f"{BASE}/start", json=wrong_result).json()["code"] == "result_digest_mismatch"
    assert db_path.exists() is False

    invalid = _input()
    invalid["periods"][1]["benefits_paid"] = "100"
    rejected = client.post(f"{BASE}/preview", json=invalid)
    assert rejected.status_code == 422
    assert rejected.json()["rows"] == []
    invalid_start = _start(invalid, {
        "input_digest": digest(invalid),
        "result_digest": digest(rejected.json()),
    }, key="health-invalid-1")
    assert client.post(f"{BASE}/start", json=invalid_start).status_code == 422
    assert db_path.exists() is False
    assert client.post(f"{BASE}/preview", content="{", headers={"content-type": "application/json"}).status_code == 400

    stored = client.post(f"{BASE}/start", json=request).json()
    changed = _input()
    changed["periods"][0]["capital_contribution"] = "6"
    changed_preview = client.post(f"{BASE}/preview", json=changed).json()
    conflict = client.post(f"{BASE}/start", json=_start(changed, changed_preview))
    assert conflict.status_code == 409
    assert conflict.json()["code"] == "idempotency_key_conflict"

    with sqlite3.connect(db_path) as connection:
        connection.execute(
            "UPDATE health_chain_results SET stored_at = ? WHERE result_id = ?",
            ("2000-01-01T00:00:00+00:00", stored["result_id"]),
        )
    url = f"{BASE}/result/{stored['result_id']}"
    assert client.get(url).json()["code"] == "stored_health_result_corrupt"
    with sqlite3.connect(db_path) as connection:
        connection.execute(
            "UPDATE health_chain_results SET stored_at = ? WHERE result_id = ?",
            (stored["stored_at"], stored["result_id"]),
        )
        connection.execute(
            "UPDATE health_chain_results SET report_json = ? WHERE result_id = ?",
            ('{"valid":true,"rows":[]}', stored["result_id"]),
        )
    assert client.get(url).json()["code"] == "stored_health_result_corrupt"
    assert client.get(f"{BASE}/results").json()["code"] == "stored_health_result_corrupt"
    for extension in ("csv", "json", "xlsx"):
        assert client.get(
            f"{url}.{extension}", headers={"If-Match": '"' + stored["result_digest"] + '"'},
        ).status_code == 409
    assert client.post(f"{BASE}/start", json=request).status_code == 409

    with sqlite3.connect(db_path) as connection:
        connection.execute(
            "UPDATE health_chain_results SET report_json = ? WHERE result_id = ?",
            ('{"valid":true,"rows":"not-a-list"}', stored["result_id"]),
        )
    assert client.get(url).json()["code"] == "stored_health_result_corrupt"


def test_start_requires_sqlite_and_timeout_never_writes(monkeypatch, tmp_path) -> None:
    module = importlib.import_module("ims.api.app")
    monkeypatch.delenv("IMS_METADATA_DB", raising=False)
    client = TestClient(module.create_app(frontend_dist=tmp_path))
    value = _input()
    preview = client.post(f"{BASE}/preview", json=value).json()
    assert client.post(f"{BASE}/start", json=_start(value, preview)).status_code == 503

    db_path = tmp_path / "metadata.sqlite"
    monkeypatch.setenv("IMS_METADATA_DB", str(db_path))
    monkeypatch.setattr(module, "HEALTH_API_TIMEOUT_SECONDS", 0.01)
    original = module.run_health_period_chain

    def wait_for_cancel(value, *, should_cancel):
        while not should_cancel():
            time.sleep(0.001)
        return original(value, should_cancel=should_cancel)

    monkeypatch.setattr(module, "run_health_period_chain", wait_for_cancel)
    client = TestClient(module.create_app(frontend_dist=tmp_path))
    timeout = client.post(f"{BASE}/start", json=_start(value, preview))
    assert timeout.status_code == 504
    assert timeout.json()["code"] == "health_calculation_timeout"
    assert db_path.exists() is False


def test_start_contract_and_missing_result_fail_without_storage(monkeypatch, tmp_path) -> None:
    module = importlib.import_module("ims.api.app")
    db_path = tmp_path / "metadata.sqlite"
    monkeypatch.setenv("IMS_METADATA_DB", str(db_path))
    client = TestClient(module.create_app(frontend_dist=tmp_path))
    value = _input()
    preview = client.post(f"{BASE}/preview", json=value).json()
    request = _start(value, preview)
    for changed, code in (
        ({**request, "schema_version": "ims.health-result-start.v0"}, "start_request_version_invalid"),
        ({**request, "idempotency_key": "short"}, "idempotency_key_invalid"),
        ({**request, "extra": True}, "start_request_fields_invalid"),
    ):
        response = client.post(f"{BASE}/start", json=changed)
        assert response.status_code == 400
        assert response.json()["code"] == code
    assert db_path.exists() is False
    assert client.get(f"{BASE}/result/not-a-result").status_code == 400
    assert client.get(f"{BASE}/result/health-{'0' * 24}").status_code == 404


def test_parallel_identical_starts_create_one_verified_record(tmp_path) -> None:
    value = _input()
    report = run_health_period_chain(value).to_dict()
    request = parse_start(_start(value, {
        "input_digest": digest(value), "result_digest": digest(report),
    }))
    db_path = tmp_path / "metadata.sqlite"
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(
            lambda _: persist_health_result(db_path, request, report), range(2),
        ))
    assert {result["result_id"] for result in results} == {results[0]["result_id"]}
    assert sorted(result["writes_performed"] for result in results) == [False, True]
    with sqlite3.connect(db_path) as connection:
        assert connection.execute("SELECT COUNT(*) FROM health_chain_results").fetchone()[0] == 1


def test_100_period_case_is_stored_and_exportable(monkeypatch, tmp_path) -> None:
    module = importlib.import_module("ims.api.app")
    monkeypatch.setenv("IMS_METADATA_DB", str(tmp_path / "metadata.sqlite"))
    client = TestClient(module.create_app(frontend_dist=tmp_path))
    value = _horizon_100()
    preview = client.post(f"{BASE}/preview", json=value)
    assert preview.status_code == 200
    started = client.post(
        f"{BASE}/start", json=_start(value, preview.json(), key="health-100-1"),
    )
    assert started.status_code == 200
    record = started.json()
    assert record["report"]["calculated_period_count"] == 100
    assert record["report"]["rows"][-1]["closing_cash"] == "1100.0000"
    url = f"{BASE}/result/{record['result_id']}"
    assert client.get(url).json()["report"] == record["report"]
    headers = {"If-Match": f'"{record["result_digest"]}"'}
    csv_rows = list(csv.DictReader(StringIO(
        client.get(f"{url}.csv", headers=headers).content.decode("utf-8"),
    )))
    assert len(csv_rows) == 100
    assert csv_rows[-1]["closing_cash"] == "1100.0000"
    assert client.get(f"{url}.json", headers=headers).json() == client.get(url).json()
    workbook = load_workbook(
        BytesIO(client.get(f"{url}.xlsx", headers=headers).content), read_only=True,
    )
    try:
        assert len(_sheet_rows(workbook["Perioden"])) == 100
        assert _sheet_rows(workbook["Perioden"])[-1]["closing_cash"] == "1100.0000"
    finally:
        workbook.close()
