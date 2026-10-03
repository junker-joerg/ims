"""Single-VU Excel uses the same checked common-market rows; text cells are safe."""
from io import BytesIO
import json
from collections.abc import Iterator

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell


def json_chunks(value: str) -> Iterator[str]:
    """Excel counts UTF-16 units; preserve astral characters and full JSON."""
    encoded, start = value.encode("utf-16-le"), 0
    while start < len(encoded):
        end = min(start + 60000, len(encoded))
        if end < len(encoded) and 0xD800 <= int.from_bytes(encoded[end - 2:end], "little") <= 0xDBFF:
            end -= 2
        yield encoded[start:end].decode("utf-16-le")
        start = end


def single_vu_workbook(result: dict, insurer_id: int) -> bytes:
    actor = next(a for a in result["source_input"]["insurers"] if a["insurer_id"] == insurer_id)
    book = Workbook(write_only=True)

    def append(sheet, values):
        cells = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=str(value) if value is not None else "")
            cell.data_type = "s"
            cells.append(cell)
        sheet.append(cells)

    for side in ("baseline", "variant"):
        sheet = book.create_sheet("Baseline" if side == "baseline" else "Variante")
        rows = [r for r in result["sides"][side]["vu_rows"] if r["insurer_id"] == insurer_id]
        fields = sorted({key for row in rows for key in row})
        append(sheet, fields)
        for row in rows:
            append(sheet, [json.dumps(row[field], ensure_ascii=False) if isinstance(row.get(field), (list, dict)) else row.get(field) for field in fields])
    trace = book.create_sheet("Kundenbuchung")
    append(trace, ["Seite", "Buchung_JSON"])
    for side in ("baseline", "variant"):
        for row in result["sides"][side]["customer_decisions"]:
            if row["insurer_id"] == insurer_id or row["previous_insurer_id"] == insurer_id:
                append(trace, [side, json.dumps(row, ensure_ascii=False)])
    source = book.create_sheet("Herkunft")
    for key, value in {"content_digest": result["content_digest"], "input_digest": result["input_digest"],
                       "contract": result["schema_version"], "insurer_id": insurer_id,
                       "insurance_group_id": actor["insurance_group_id"], "name": actor["name"],
                       "units": "model_currency", "limits": json.dumps(result["limits"]) }.items():
        append(source, [key, value])
    # Reproducible complete market context, chunked below Excel cell limits.
    encoded = json.dumps(result["source_input"], ensure_ascii=False, allow_nan=False)
    for chunk in json_chunks(encoded):
        append(source, ["source_input_json_chunk", chunk])
    if "source_bundle" in result:
        reference = result["reference"]
        group = next((g for g in reference["groups"] if g["insurer_id"] == insurer_id), None)
        if group is None:
            group = {"insurer_id": insurer_id, "name": actor["name"], "entity_rows": [], "synthetic_extra_actor": True}
        facts = book.create_sheet("BaFin-Quellenwerte")
        append(facts, ["Verdiente Bruttobeiträge, Mio. EUR; einschließlich Ausland und übernommener Rückversicherung. Keine konzerninterne Eliminierung; kein vollständiger deutscher Direktmarkt."])
        append(facts, ["Redaktionelle Gruppenzuordnung, noch keine vollständige unabhängige Kontrollprüfung."])
        if group.get("synthetic_extra_actor"):
            append(facts, ["Fiktiver Zusatzanbieter: keine BaFin-Quellengruppe; nachstehende Dateien belegen nur den 40er-Ausgangsmarkt."])
        fields = ["workbook_row", "entity_source_name", "source_sector", "raw_value", "source_unit", "premium_million_eur", "value_status", "source_id", "source_sheet", "source_cell"]
        append(facts, fields)
        for row in result["source_bundle"]["source_catalog"]["entities"]:
            if row["workbook_row"] in group["entity_rows"]:
                append(facts, [row[f] for f in fields])
        for entry in result["source_bundle"]["source_catalog"]["source_files"]:
            append(facts, [entry["id"], entry["url"], entry["sha256"]])
        assumptions = book.create_sheet("Workshop-Annahmen")
        for key, value in {"source_scope": reference["source_scope"], "selected_group": group,
                           "reference_universe_million_eur": reference["universe_total_million_eur"],
                           "selected_total_million_eur": reference["selected_total_million_eur"],
                           "rest": reference["rest"], "workshop": result["source_bundle"]["workshop"],
                           "overrides": result["source_bundle"]["overrides"],
                           "model_binding": reference["model_binding"], "model_binding_note": reference["model_binding_note"]}.items():
            append(assumptions, [key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
        bundle = json.dumps(result["source_bundle"], ensure_ascii=False, allow_nan=False)
        for chunk in json_chunks(bundle):
            append(source, ["source_bundle_json_chunk", chunk])
    output = BytesIO()
    if "shock_bundle" in result:
        for title, name in (("ICT-Prozesse", "ict_process_rows"), ("ICT-Abhängigkeiten", "ict_dependency_rows"),
                            ("ICT-Ressourcen", "ict_resource_rows"), ("ICT-Ereignisse", "ict_event_rows"),
                            ("Kosten-einmal", "cost_rows"), ("Lebens-Anträge", "life_demand_rows"),
                            ("Neue-Lebenspolicen", "life_contract_rows"), ("Wirksame-Wechsel", "switch_process_rows"),
                            ("Horizont-Rückstand", "terminal_process_rows")):
            sheet = book.create_sheet(title)
            append(sheet, ["Seite", "Einmalige Buchung / Vorgang / Abhängigkeit; Modellannahmen, keine recherchierten Firmenprofile"])
            for side in ("baseline", "variant"):
                for row in result["sides"][side][name]:
                    if "insurer_id" not in row or row["insurer_id"] == insurer_id or name == "switch_process_rows" and row.get("previous_insurer_id") == insurer_id:
                        append(sheet, [side, json.dumps(row, ensure_ascii=False)])
        encoded = json.dumps(result["shock_bundle"], ensure_ascii=False, allow_nan=False)
        for chunk in json_chunks(encoded):
            append(source, ["shock_bundle_json_chunk", chunk])
    book.save(output)
    return output.getvalue()
