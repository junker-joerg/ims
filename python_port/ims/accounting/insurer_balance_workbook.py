"""In-memory XLSX presentation of a verified insurer model balance."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

from ims.accounting.insurer_balance import InsurerBalanceReport
from ims.accounting.model_balance_contract import BALANCE_FIELDS


_LABELS = {
    "period": "Periode",
    "opening_cash": "Cash Anfang",
    "opening_claim_liability": "Schadenverbindlichkeit Anfang",
    "opening_equity": "Eigenkapital Anfang",
    "premium_income": "Praemieneinnahme",
    "investment_income": "Zinsertrag",
    "claims_incurred": "Angefallene Schaeden",
    "claims_paid": "Bezahlte Schaeden",
    "operating_expense": "Laufender Aufwand",
    "capital_contribution": "Kapitalzufuehrung",
    "capital_distribution": "Kapitalausschuettung",
    "period_profit": "Periodenergebnis",
    "closing_cash": "Cash Schluss",
    "closing_claim_liability": "Schadenverbindlichkeit Schluss",
    "closing_equity": "Eigenkapital Schluss",
}
_COLUMNS = ("period", *(field.field_id for field in BALANCE_FIELDS))


def build_insurer_balance_workbook(report: InsurerBalanceReport) -> bytes:
    """Export exact canonical amounts as text cells; never round via Excel floats."""

    payload = report.to_dict()
    if not payload["valid"] or not payload["content_digest"]:
        raise ValueError("Nur eine vollstaendig gueltige Versichererbilanz ist exportierbar")

    workbook = Workbook(write_only=True)
    for name, rows in (
        ("Gesamt", payload["total_rows"]),
        ("Kfz", payload["sectors"][0]["rows"]),
        ("Sach-Haftpflicht", payload["sectors"][1]["rows"]),
    ):
        sheet = workbook.create_sheet(name)
        sheet.freeze_panes = "B2"
        sheet.append([_LABELS[column] for column in _COLUMNS])
        for row in rows:
            cells = [row["period"]]
            for column in _COLUMNS[1:]:
                cell = WriteOnlyCell(sheet, value=row[column])
                cell.data_type = "s"
                cells.append(cell)
            sheet.append(cells)

    provenance = workbook.create_sheet("Herkunft")
    for label, value in (
        ("Versicherer-ID", payload["insurer_id"]),
        ("Perioden", payload["period_count"]),
        ("Quelle", "Explizite Szenariowerte"),
        ("Historische Spartenbindung", "Offen"),
        ("Modellgrenze", "Keine gesetzliche Bilanz oder Solvency-II-Meldung"),
        ("Betragszellen", "Exakte Dezimalstrings als Text; keine Excel-Rundung"),
        ("Ergebnis-Digest", payload["content_digest"]),
        ("Schema", payload["schema_version"]),
    ):
        provenance.append([label, value])

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
