"""In-memory workbook for the model-only capital workshop result."""

from __future__ import annotations

import hashlib
import json
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill


_SECTOR_LABELS = {
    "motor": "Kfz",
    "property_liability": "Sach-Haftpflicht",
    "life": "Leben",
    "health": "Kranken",
}
_RISK_LABELS = {
    "asset_market_value": "Marktwert Aktiva",
    "non_life_claim_obligation": "Schadenverpflichtung",
    "life_obligation": "Lebensverpflichtung",
    "health_benefit_obligation": "Krankenleistung",
    "counterparty_model_loss": "Gegenpartei",
    "operational_model_loss": "Betriebsereignis",
}
_REGULATORY_LABELS = {
    "eligible_own_funds": "Anrechenbare Eigenmittel",
    "scr": "SCR",
    "mcr": "MCR",
    "scr_coverage_ratio": "SCR-Bedeckungsquote",
    "mcr_coverage_ratio": "MCR-Bedeckungsquote",
}


def _header(sheet, labels: list[str]) -> None:
    cells = []
    for label in labels:
        cell = WriteOnlyCell(sheet, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="245C45")
        cells.append(cell)
    sheet.append(cells)


def build_solvency_capital_workbook(
    input_payload: dict, readiness: dict, balance: dict, aggregation: dict,
) -> bytes:
    """Export exact model amounts; regulatory metrics remain blank and blocked."""

    if not all(report.get("valid") and report.get("content_digest") for report in (
        readiness, balance, aggregation,
    )):
        raise ValueError("Nur vollstaendig gueltige Kapital-Modellergebnisse sind exportierbar")
    input_digest = hashlib.sha256(
        json.dumps(input_payload, sort_keys=True, separators=(",", ":"), ensure_ascii=True).encode("utf-8")
    ).hexdigest()
    if readiness.get("input_digest") != input_digest:
        raise ValueError("Kapital-Eingang und Ergebnis-Digest stimmen nicht ueberein")
    if (
        readiness["source_model_balance_content_digest"] != balance["content_digest"]
        or readiness["source_risk_aggregation_content_digest"] != aggregation["content_digest"]
        or aggregation["source_model_balance_content_digest"] != balance["content_digest"]
    ):
        raise ValueError("Kapital-Modellergebnisse haben widerspruechliche Quellen-Digests")
    regulatory = readiness["regulatory_metrics"]
    if regulatory.get("status") != "blocked_missing_regulatory_basis" or any(
        regulatory.get(key) is not None for key in _REGULATORY_LABELS
    ):
        raise ValueError("Regulatorische Kennzahlen muessen gesperrt bleiben")

    workbook = Workbook(write_only=True)
    model = workbook.create_sheet("Modellwirkung")
    model.freeze_panes = "B2"
    _header(model, ["Ebene", "Kennzahl", "Wert", "Einheit oder Bedeutung"])
    management = readiness["management_evaluation"]
    totals = aggregation["totals"]
    for level, label, value, meaning in (
        ("Modell", "Eigenmittel-Proxy", management["model_own_funds_proxy"], "Modellwaehrung, nicht EUR"),
        ("Modell", "Brutto-Stressverlust", totals["gross_model_stress_loss"], "Szenarioannahme"),
        ("Modell", "Angerechneter Modellpuffer", totals["model_buffer_applied"], "Eigenstaendig deklariert"),
        ("Modell", "Netto-Stressverlust", management["net_model_stress_loss"], "Kein SCR"),
        ("Modell", "Verbleibender Eigenkapital-Proxy", management["remaining_model_equity_proxy"], "Keine Bilanzbuchung"),
        ("Workshop", "Maximaler Nettoverlust", management["max_net_model_stress_loss"], "Inklusive Grenze"),
        ("Workshop", "Minimaler Restproxy", management["min_remaining_model_equity_proxy"], "Inklusive Grenze"),
        ("Workshop", "Verlustgrenze eingehalten", "Ja" if management["loss_limit_met"] else "Nein", "Nicht regulatorisch"),
        ("Workshop", "Restproxy-Grenze eingehalten", "Ja" if management["equity_floor_met"] else "Nein", "Nicht regulatorisch"),
    ):
        model.append([level, label, value, meaning])

    sectors = workbook.create_sheet("Sparten")
    sectors.freeze_panes = "B2"
    _header(sectors, ["Sparte", "Aktiva angepasst", "Verpflichtungen angepasst", "Eigenmittel-Proxy"])
    for row in balance["sector_rows"]:
        sectors.append([
            _SECTOR_LABELS[row["sector_id"]], row["adjusted_model_assets"],
            row["adjusted_model_liabilities"], row["model_own_funds_proxy"],
        ])
    total = balance["total_row"]
    sectors.append(["Gesamt", total["adjusted_model_assets"], total["adjusted_model_liabilities"], total["model_own_funds_proxy"]])

    risks = workbook.create_sheet("Risiken")
    risks.freeze_panes = "B2"
    _header(risks, ["Modellkomponente", "Verlust", "Faktorladung"])
    for row in aggregation["component_rows"]:
        risks.append([
            _RISK_LABELS[row["risk_kind"]], row["model_loss_amount"],
            row["common_factor_loading"],
        ])
    risks.append(["Summe vor Aggregation", totals["unadjusted_sum_of_components"], None])
    risks.append(["Diversifikations-Proxy", totals["scenario_diversification_proxy"], None])

    assumptions = workbook.create_sheet("Annahmen")
    assumptions.freeze_panes = "B2"
    _header(assumptions, ["Bereich", "Position", "Parameter", "Wert"])
    aggregation_input = input_payload["solvency_risk_aggregation_input"]
    module_input = aggregation_input["solvency_risk_modules_input"]
    shock_input = module_input["solvency_scenario_shocks_input"]
    model_input = shock_input["solvency_model_balance_input"]
    for row in model_input["adjustments"]:
        for field in ("asset_delta", "liability_delta"):
            assumptions.append(["Modellbewertung", row["sector_id"], field, row[field]])
    for row in shock_input["exposures"]:
        assumptions.append(["Teilposition", row["exposure_id"], "base_amount", row["base_amount"]])
    for row in module_input["module_parameters"]:
        assumptions.append(["Modulsatz", row["exposure_id"], "stress_rate", row["stress_rate"]])
    for row in aggregation_input["counterparty_cases"]:
        assumptions.append(["Gegenpartei", row["exposure_id"], "loss_rate", row["loss_rate"]])
    for row in aggregation_input["operational_events"]:
        assumptions.append(["Betriebsereignis", row["event_id"], "loss_amount", row["loss_amount"]])
    for kind, loading in aggregation_input["aggregation_assumptions"]["factor_loadings"].items():
        assumptions.append(["Faktorladung", kind, "loading", loading])
    for field in ("capacity_amount", "applied_amount"):
        assumptions.append(["Modellpuffer", "workshop_buffer", field, aggregation_input["loss_absorption"][field]])
    for field in ("max_net_model_stress_loss", "min_remaining_model_equity_proxy"):
        assumptions.append(["Workshop-Grenze", "management", field, input_payload["management_limits"][field]])

    blocked = workbook.create_sheet("Regulatorik")
    blocked.freeze_panes = "B2"
    _header(blocked, ["Kennzahl", "Wert", "Status"])
    for key, label in _REGULATORY_LABELS.items():
        blocked.append([label, None, "Nicht berechnet: regulatorische Grundlage fehlt"])
    blocked.append([])
    _header(blocked, ["Offene Voraussetzung", "Status", "Hinweis"])
    for row in readiness["readiness_rows"]:
        blocked.append([row["requirement_code"], row["status"], row["message"]])

    provenance = workbook.create_sheet("Herkunft")
    _header(provenance, ["Merkmal", "Wert"])
    for label, value in (
        ("Versicherer-ID", readiness["insurer_id"]),
        ("Szenario", readiness["scenario_id"]),
        ("Variante", readiness["variant_id"]),
        ("Modellperiode", readiness["checkpoint"]["model_period"]),
        ("Szenario-Referenztag", readiness["checkpoint"]["reference_date"]),
        ("Einheit", readiness["amount_unit"]),
        ("Eingang-Digest", readiness["input_digest"]),
        ("Vier-Sparten-Bilanz-Digest", balance["source_content_digest"]),
        ("PR172-Modellbilanz-Digest", balance["content_digest"]),
        ("PR175-Stress-Digest", aggregation["content_digest"]),
        ("PR176-Kapitalansicht-Digest", readiness["content_digest"]),
        ("Abgrenzung", "Workshop-Modell, kein SCR/MCR, keine Compliance- oder historische Vollgleichheitsaussage"),
        ("Betragsdarstellung", "Exakte Dezimaltexte als Excel-Textzellen"),
    ):
        provenance.append([label, value])

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
