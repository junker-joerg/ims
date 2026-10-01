import csv
from io import BytesIO, StringIO
import json
from threading import BoundedSemaphore

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from starlette.applications import Starlette
from starlette.concurrency import run_in_threadpool
from starlette.responses import JSONResponse, Response
from starlette.routing import Route

from ims.accounting.management_case import calculate, declare_sources, workshop_case

MAX_INPUT_BYTES = 4 * 1024 * 1024


async def request_json(request):
    """Bound every management input, including preset requests, before decoding."""
    chunks, size = [], 0
    async for chunk in request.stream():
        size += len(chunk)
        if size > MAX_INPUT_BYTES:
            raise ValueError("Management-Eingabe überschreitet 4 MiB")
        chunks.append(chunk)
    try:
        return json.loads(b"".join(chunks))
    except (ValueError, UnicodeDecodeError) as exc:
        raise ValueError("Gültiges JSON erforderlich") from exc


def flat_rows(result: dict) -> list[dict]:
    return [{"content_digest": result["content_digest"], "scenario_id": result["scenario_id"], "insurer_id": result["insurer_id"], "variant_id": side, "sector_id": sector["sector_id"], **row}
            for side in ("baseline", "variant") for sector in [*result["sides"][side]["sectors"], {"sector_id": "total", "rows": result["sides"][side]["total_rows"]}] for row in sector["rows"]]


def exports(result: dict, extension: str) -> tuple[bytes, str]:
    if extension == "json": return json.dumps(result, ensure_ascii=False, allow_nan=False).encode("utf-8"), "application/json"
    rows = flat_rows(result)
    fields = list(rows[0])
    if extension == "csv":
        output = StringIO(newline="")
        writer = csv.DictWriter(output, fieldnames=fields)
        writer.writeheader(); writer.writerows(rows)
        return output.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8"
    book = Workbook(write_only=True)
    def append(sheet, values):
        cells = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=str(value)); cell.data_type = "s"; cells.append(cell)
        sheet.append(cells)
    for side in ("baseline", "variant"):
        sheet = book.create_sheet(side)
        append(sheet, fields)
        for row in rows:
            if row["variant_id"] == side: append(sheet, [row[field] for field in fields])
    source = book.create_sheet("Herkunft")
    append(source, ["content_digest", result["content_digest"]])
    append(source, ["coupling_scope", result["coupling_scope"]])
    append(source, ["regulatory_metrics", json.dumps(result["regulatory_metrics"])])
    for side in ("baseline", "variant"):
        append(source, [side, json.dumps(result["source_contracts"][side], ensure_ascii=False)])
    # Complete input, chunked below Excel's 32767-character cell boundary.
    full_input = json.dumps(result["source_input"], ensure_ascii=False, allow_nan=False)
    for index in range(0, len(full_input), 30000): append(source, ["source_input_json_chunk", full_input[index:index + 30000]])
    output = BytesIO(); book.save(output)
    return output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def create_management_app() -> Starlette:
    gate = BoundedSemaphore(1)
    headers = {"Cache-Control": "no-store"}
    def failed(message): return {"valid": False, "issues": [{"path": "$", "code": "request_invalid", "message": message}], "content_digest": None, "sides": {}, "partial_result_returned": False, "writes_performed": False}

    async def preset(request):
        try:
            value = await request_json(request)
            if not isinstance(value, dict) or set(value) != {"case_id", "period_count", "insurer_id"}: raise ValueError("Fall, Horizont und VU-ID ausdrücklich angeben")
            source = workshop_case(value["case_id"], value["period_count"], value["insurer_id"])
            return JSONResponse({"valid": True, "source_input": source, "writes_performed": False}, headers=headers)
        except (ValueError, TypeError) as exc: return JSONResponse(failed(str(exc)), status_code=422, headers=headers)

    async def action(request):
        try: value = await request_json(request)
        except ValueError as exc: return JSONResponse(failed(str(exc)), status_code=422, headers=headers)
        if not gate.acquire(blocking=False): return JSONResponse(failed("Eine gemeinsame Rechnung läuft bereits"), status_code=409, headers=headers)
        def run():
            try: return declare_sources(value) if request.url.path.endswith("/declare-sources") else calculate(value)
            finally: gate.release()
        result = await run_in_threadpool(run)
        if not result["valid"]: return JSONResponse(result, status_code=422, headers=headers)
        etag = '"' + result["content_digest"] + '"'
        extension = request.path_params.get("extension")
        reply_headers = {**headers, "ETag": etag}
        if not extension: return JSONResponse(result, headers=reply_headers)
        if extension not in ("json", "csv", "xlsx"): return JSONResponse(failed("Unbekanntes Exportformat"), status_code=404, headers=headers)
        if request.headers.get("if-match") != etag: return JSONResponse(failed("Export benötigt denselben unveränderten Ergebnisnachweis"), status_code=412, headers=headers)
        content, media = await run_in_threadpool(exports, result, extension)
        return Response(content, media_type=media, headers={**reply_headers, "Content-Disposition": f'attachment; filename="IMS-Management-{result["content_digest"][:12]}.{extension}"'})

    return Starlette(routes=[Route("/workshop-case", preset, methods=["POST"]), Route("/declare-sources", action, methods=["POST"]), Route("/calculate", action, methods=["POST"]), Route("/export.{extension}", action, methods=["POST"])])
