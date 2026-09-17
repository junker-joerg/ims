"""Exact CSV, JSON and XLSX exports from a verified stored health result."""

from __future__ import annotations

import csv
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

from ims.api.health_result_delivery import canonical_json


_PROVENANCE = (
    ("result_id", "Ergebnis-ID"),
    ("input_digest", "Eingabe-Digest"),
    ("result_digest", "Ergebnis-Digest"),
    ("record_digest", "Datensatz-Digest"),
)


def _verified(record: dict) -> None:
    if record.get("stored_result_verified") is not True:
        raise ValueError("Nur gepruefte gespeicherte Krankenergebnisse exportieren")


def _row_provenance(record: dict) -> dict[str, object]:
    report = record["report"]
    return {
        **{field: record[field] for field, _ in _PROVENANCE},
        "insurer_id": report["insurer_id"],
        "scenario_id": report["scenario_id"],
        "variant_id": report["variant_id"],
        "period_count": report["calculated_period_count"],
    }


def build_health_result_csv(record: dict) -> bytes:
    _verified(record)
    rows = record["report"]["rows"]
    provenance = _row_provenance(record)
    fields = [
        *provenance, "period",
        *sorted(key for key in rows[0] if key != "period"),
    ]
    stream = StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=fields, lineterminator="\n")
    writer.writeheader()
    for row in rows:
        writer.writerow({**provenance, **row})
    return stream.getvalue().encode("utf-8")


def build_health_result_json(record: dict) -> bytes:
    _verified(record)
    return canonical_json(record).encode("ascii")


def _text_cell(sheet, value: object) -> object:
    if type(value) is str:
        cell = WriteOnlyCell(sheet, value=value)
        cell.data_type = "s"
        return cell
    return value


def build_health_result_workbook(record: dict) -> bytes:
    _verified(record)
    report = record["report"]
    rows = report["rows"]
    workbook = Workbook(write_only=True)
    periods = workbook.create_sheet("Perioden")
    periods.freeze_panes = "B2"
    fields = ["period", *sorted(key for key in rows[0] if key != "period")]
    periods.append(fields)
    for row in rows:
        periods.append([_text_cell(periods, row[field]) for field in fields])

    provenance = workbook.create_sheet("Herkunft")
    source = record["health_chain_input"]["sources"]
    for label, value in (
        *((label, record[field]) for field, label in _PROVENANCE),
        ("Versicherer-ID", report["insurer_id"]),
        ("Szenario", report["scenario_id"]),
        ("Variante", report["variant_id"]),
        ("Perioden", report["calculated_period_count"]),
        ("Quellen-Schema", source["schema_version"]),
        ("Neugeschaeft", source["new_business"]["mode"]),
        ("Abgang", source["exits"]["mode"]),
        ("Beitrag", source["pricing"]["mode"]),
        ("Leistungsanfall", source["benefits"]["mode"]),
        ("Betragszellen", "Exakte Dezimalstrings als Text"),
        ("Modell", "IMS 2.x Kranken; keine historische Vollgleichheit"),
        ("Bilanz", "Modellbilanz, keine gesetzliche oder Solvency-II-Bilanz"),
    ):
        provenance.append([label, _text_cell(provenance, value)])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
