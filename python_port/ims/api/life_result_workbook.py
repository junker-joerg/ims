"""Exact-text workbook for a verified stored IMS 2.x life result."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell


_ROW_NESTED = {
    "opening_cohorts": "Kohorten Anfang",
    "closing_cohorts": "Kohorten Schluss",
    "cohort_movements": "Kohortenbewegung",
    "new_business_cohorts": "Neugeschaeft Kohorten",
    "opening_policies": "Policen Anfang",
    "closing_policies": "Policen Schluss",
    "policy_movements": "Policenbewegung",
    "new_policy_issues": "Neupolicen",
}


def _cell(sheet, value: object) -> object:
    if type(value) is str:
        cell = WriteOnlyCell(sheet, value=value)
        cell.data_type = "s"
        return cell
    return value


def _table(workbook: Workbook, name: str, values: list[dict], *, period: bool = False) -> None:
    sheet = workbook.create_sheet(name)
    sheet.freeze_panes = "B2"
    fields = sorted({key for value in values for key in value})
    if period:
        fields = ["period", *[field for field in fields if field != "period"]]
    sheet.append(fields)
    for value in values:
        sheet.append([_cell(sheet, value.get(field)) for field in fields])


def build_life_result_workbook(record: dict) -> bytes:
    """Keep every exported amount byte-equivalent to the stored JSON string."""
    if record.get("stored_result_verified") is not True:
        raise ValueError("Nur gepruefte gespeicherte Lebensergebnisse exportieren")
    report = record["report"]
    workbook = Workbook(write_only=True)
    rows = report["rows"]
    _table(workbook, "Perioden", [
        {key: value for key, value in row.items() if key not in _ROW_NESTED}
        for row in rows
    ], period=True)
    _table(workbook, "Quellen", [
        {
            "period": source["period"],
            "investment_mode": source["investment_mode"],
            "mortality_mode": source["mortality_mode"],
            "opening_backing_assets": source["opening_backing_assets"],
            "opening_active_policies": source["opening_active_policies"],
            "investment_rate_per_period": source["investment_rate_per_period"],
            "mortality_rate_per_period": source["mortality_rate_per_period"],
            "investment_result": source["investment_result"],
            "death_policy_ids": ", ".join(source["death_policy_ids"]),
            "death_count": source["death_count"],
        }
        for source in report["resolved_sources"]
    ], period=True)
    _table(workbook, "Todesfaelle Kohorten", [
        {"period": source["period"], **item}
        for source in report["resolved_sources"]
        for item in source["deaths_by_cohort"]
    ], period=True)
    for field, name in _ROW_NESTED.items():
        _table(workbook, name, [
            {"period": row["period"], **item}
            for row in rows for item in row[field]
        ], period=True)
    provenance = workbook.create_sheet("Herkunft")
    for label, value in (
        ("Ergebnis-ID", record["result_id"]),
        ("Eingabe-Digest", record["input_digest"]),
        ("Ergebnis-Digest", record["result_digest"]),
        ("Datensatz-Digest", record["record_digest"]),
        ("Versicherer-ID", report["insurer_id"]),
        ("Perioden", report["calculated_period_count"]),
        ("Modell", "IMS 2.x Leben; keine historische Vollgleichheit"),
        ("Bilanz", "Modellbilanz, keine gesetzliche oder Solvency-II-Bilanz"),
        ("Betragszellen", "Exakte Dezimalstrings als Text"),
        ("Ergebnis-Schema", report["schema_version"]),
    ):
        provenance.append([label, _cell(provenance, value)])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
