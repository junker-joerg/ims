"""Read-only audit of the supplied AP6 BaFin reference, never a model importer.

Recompute in Decimal, compare each quoted value with pinned BaFin originals,
and retain the workbook's editorial group assignments as unverified assertions.
No Excel formulas, URLs or document instructions are executed. No source is saved.
"""
from __future__ import annotations

import argparse
from collections import Counter
from decimal import Decimal
from hashlib import sha256
import json
from pathlib import Path
import warnings

import openpyxl


ROOT = Path(__file__).resolve().parents[2]
SOURCE_REGISTER = ROOT / "docs/research/ims_ap6_sources_2026_10.json"
ZERO = Decimal(0)
# Comparison tolerance only for Excel's binary floating-point caches, not ranking.
CACHE_TOLERANCE = Decimal("0.0000001")
LAYOUT = {
    "160": ("AP6-S02", "Leben", "Mio. EUR", Decimal(1), 80, "C11"),
    "460": ("AP6-S03", "Kranken", "Tsd. EUR", Decimal("0.001"), 43, "C11"),
    "560": ("AP6-S04", "Schaden/Unfall", "Mio. EUR", Decimal(1), 203, "C12"),
}


def numeric(value: object) -> Decimal | None:
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        number = Decimal(str(value))
        if number.is_finite() and number >= 0:
            return number
    return None


def premium(value: object, factor: Decimal) -> tuple[Decimal | None, Decimal | None, Decimal | None, str]:
    """BaFin PDF p.2: '-' exact zero; numeric 0 below the stated unit.

    Other numeric values have a conservative one-source-unit envelope. Half-unit
    rounding is not assumed. Bounds apply only to the quoted earned measure.
    """
    if value == "-":
        return ZERO, ZERO, ZERO, "source_exact_zero"
    number = numeric(value)
    if number is None:
        return None, None, None, "unknown"
    amount = number * factor
    if number == 0:
        return amount, ZERO, factor, "source_below_unit"
    return amount, max(ZERO, amount - factor), amount + factor, "source_rounded"


def aggregate(rows: list[dict]) -> list[dict]:
    groups: dict[str, dict] = {}
    for row in rows:
        name = row["group_assertion"]
        group = groups.setdefault(name, {
            "group_assertion": name, "entity_rows": [], "total": ZERO,
            "lower": ZERO, "upper": ZERO, "unknown_count": 0,
            "sectors": {sector: ZERO for _, sector, *_ in LAYOUT.values()},
            "sector_entity_count": {sector: 0 for _, sector, *_ in LAYOUT.values()},
        })
        group["entity_rows"].append(row["workbook_row"])
        group["sector_entity_count"][row["source_sector"]] += 1
        value = row["premium_million_eur"]
        if value is None:
            group["unknown_count"] += 1
        else:
            group["total"] += Decimal(value)
            group["sectors"][row["source_sector"]] += Decimal(value)
            group["lower"] += Decimal(row["lower_million_eur"])
            group["upper"] += Decimal(row["upper_million_eur"])
    ranked = sorted(groups.values(), key=lambda g: (-g["total"], g["group_assertion"]))
    for rank, group in enumerate(ranked, 1):
        group["rank_of_available_values"] = rank
        group["entity_count"] = len(group["entity_rows"])
        if group["unknown_count"]:
            group["upper"] = None  # Missing values cannot be bounded by zero.
        group["group_membership_verified"] = False
        group["sector_note"] = "Zero with no source entity is an empty sum in this source universe, not proven inactivity."
    return ranked


def boundary(ranked: list[dict], count: int = 40) -> dict:
    if len(ranked) <= count:
        raise ValueError("Die Rangbasis muss auch nicht ausgewählte Gruppen enthalten.")
    selected, others = ranked[:count], ranked[count:]
    last, next_group = selected[-1], others[0]
    complete = not any(g["unknown_count"] for g in ranked)
    minimum = min(g["lower"] for g in selected)
    maximum = max((g["upper"] for g in others if g["upper"] is not None), default=ZERO)
    return {
        "rank_40_group_assertion": last["group_assertion"],
        "rank_40_million_eur": last["total"],
        "rank_41_group_assertion": next_group["group_assertion"],
        "rank_41_million_eur": next_group["total"],
        "gap_million_eur": last["total"] - next_group["total"],
        "minimum_selected_lower_million_eur": minimum,
        "maximum_other_upper_million_eur": maximum if complete else None,
        "separated_within_workbook_assignments_and_source_units": complete and minimum > maximum,
        "german_market_boundary_verified": False,
    }


def json_value(value: object) -> object:
    if isinstance(value, Decimal):
        return format(value, "f")
    raise TypeError(type(value).__name__)


def audit(workbook_path: Path, source_dir: Path, register_path: Path = SOURCE_REGISTER) -> dict:
    raw = workbook_path.read_bytes()
    sources = {s["id"]: s for s in json.loads(register_path.read_text(encoding="utf-8"))["sources"]}
    originals, provenance = {}, []
    for sheet, (source_id, sector, unit, factor, expected_count, total_cell) in LAYOUT.items():
        source = sources[source_id]
        path = source_dir / Path(source["local_download"]).name
        digest = sha256(path.read_bytes()).hexdigest()
        if digest != source["sha256"]:
            raise ValueError(f"SHA-256 stimmt nicht: {source_id}")
        originals[sheet] = openpyxl.load_workbook(path, data_only=True)[sheet]
        provenance.append({"id": source_id, "url": source["url"], "sha256": digest,
                           "sheet": sheet, "unit": unit, "expected_entity_count": expected_count})
    # Unsupported visual extensions are irrelevant to reading; originals are never saved.
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
        formulas = openpyxl.load_workbook(workbook_path, data_only=False)
        cached = openpyxl.load_workbook(workbook_path, data_only=True)
    if formulas.sheetnames != ["Top 40", "Einzelgesellschaften", "Methodik"]:
        raise ValueError("Nicht unterstützte Arbeitsmappenstruktur.")
    detail = formulas["Einzelgesellschaften"]
    if detail.max_row != 330:
        raise ValueError("Erwartet werden 326 Datenzeilen (5:330).")
    errors, rows, seen, checked_formulas = [], [], set(), set()

    def check(condition: bool, message: str) -> None:
        if not condition:
            errors.append(message)

    def check_formula(sheet: str, cell: str, expected: str, result: Decimal | None) -> None:
        checked_formulas.add((sheet, cell))
        check(formulas[sheet][cell].value == expected, f"Formel {sheet}!{cell} weicht ab")
        actual = cached[sheet][cell].value
        if result is None:
            check(actual in (None, ""), f"Cache {sheet}!{cell} ist nicht leer")
        else:
            number = numeric(actual)
            check(number is not None and abs(number - result) <= CACHE_TOLERANCE,
                  f"Cache {sheet}!{cell} weicht von unabhängiger Rechnung ab")

    for index in range(5, 331):
        group, name, sector, value, unit, raw_factor, _, sheet, cell, url = [c.value for c in detail[index]]
        if sheet not in LAYOUT:
            raise ValueError(f"Unbekannte Quelltabelle in Zeile {index}")
        source_id, expected_sector, expected_unit, factor, _, _ = LAYOUT[sheet]
        source = originals[sheet]
        check(isinstance(group, str) and bool(group.strip()), f"Gruppe fehlt in Zeile {index}")
        check(sector == expected_sector and unit == expected_unit and numeric(raw_factor) == factor,
              f"Sparte/Einheit/Faktor falsch in Zeile {index}")
        check(url == sources[source_id]["url"], f"Quell-URL falsch in Zeile {index}")
        if not isinstance(cell, str) or not cell.startswith("C") or not cell[1:].isdigit():
            raise ValueError(f"Ungültiger Quellzellbezug in Zeile {index}")
        source_row = int(cell[1:])
        check(source[cell].value == value, f"Quellbetrag falsch in Zeile {index}")
        check(source.cell(source_row, 2).value == name, f"Quellname falsch in Zeile {index}")
        key = (sheet, cell)
        check(key not in seen, f"Doppelter Quellzellbezug in Zeile {index}")
        seen.add(key)
        amount, lower, upper, status = premium(value, factor)
        cache_amount = numeric(value)
        check_formula("Einzelgesellschaften", f"G{index}",
                      f'=IF(ISNUMBER(D{index}),D{index}*F{index},"")',
                      cache_amount * factor if cache_amount is not None else None)
        rows.append({
            "workbook_row": index, "group_assertion": group, "entity_source_name": name,
            "source_sector": sector, "raw_value": value, "source_unit": unit,
            "source_id": source_id, "source_sheet": sheet, "source_cell": cell,
            "premium_million_eur": None if amount is None else str(amount),
            "lower_million_eur": None if lower is None else str(lower),
            "upper_million_eur": None if upper is None else str(upper), "value_status": status,
        })
    counts = Counter(row["source_sheet"] for row in rows)
    reconciliations = []
    for sheet, (_, sector, _, factor, expected_count, total_cell) in LAYOUT.items():
        source = originals[sheet]
        cells = {f"C{r}" for r in range(1, source.max_row + 1)
                 if isinstance(source.cell(r, 1).value, int) and isinstance(source.cell(r, 2).value, str)}
        check(counts[sheet] == expected_count == len(cells), f"Anzahl in Tabelle {sheet} falsch")
        check({cell for s, cell in seen if s == sheet} == cells, f"Quelluniversum {sheet} unvollständig")
        amount = sum((Decimal(row["premium_million_eur"]) for row in rows
                      if row["source_sheet"] == sheet and row["premium_million_eur"] is not None), ZERO)
        branch_total = Decimal(str(source[total_cell].value)) * factor
        reconciliations.append({"source_sheet": sheet, "source_sector": sector,
                                "entity_sum_million_eur": amount, "branch_sum_million_eur": branch_total,
                                "difference_million_eur": amount - branch_total})
    ranked = aggregate(rows)
    by_group = {g["group_assertion"]: g for g in ranked}
    displayed = []
    for r in range(7, 47):
        name = formulas["Top 40"][f"B{r}"].value
        if name not in by_group:
            raise ValueError(f"Unbekannte Gruppe in Top 40!B{r}")
        displayed.append(name)
        check(formulas["Top 40"][f"A{r}"].value == r - 6, f"Statischer Rang A{r} falsch")
        group = by_group[name]
        check_formula("Top 40", f"C{r}", f"=SUM(D{r}:F{r})", group["total"])
        for col, sector in zip("DEF", ("Leben", "Kranken", "Schaden/Unfall")):
            check_formula("Top 40", f"{col}{r}",
                          f'=SUMIFS(\'Einzelgesellschaften\'!$G$5:$G$330,\'Einzelgesellschaften\'!$A$5:$A$330,$B{r},\'Einzelgesellschaften\'!$C$5:$C$330,"{sector}")',
                          group["sectors"][sector])
        for col, numerator, sector in zip("GHI", "DEF", ("Leben", "Kranken", "Schaden/Unfall")):
            check_formula("Top 40", f"{col}{r}", f"={numerator}{r}/C{r}",
                          group["sectors"][sector] / group["total"] if group["total"] else None)
        check_formula("Top 40", f"J{r}", f"=COUNTIF('Einzelgesellschaften'!$A$5:$A$330,B{r})",
                      Decimal(group["entity_count"]))
    check(displayed == [g["group_assertion"] for g in ranked[:40]], "Statische Top 40 entsprechen nicht der unabhängigen Sortierung")
    totals = {"C": sum((by_group[name]["total"] for name in displayed), ZERO),
              "J": Decimal(sum(by_group[name]["entity_count"] for name in displayed))}
    for col, sector in zip("DEF", ("Leben", "Kranken", "Schaden/Unfall")):
        totals[col] = sum((by_group[name]["sectors"][sector] for name in displayed), ZERO)
    for col, result in totals.items():
        check_formula("Top 40", f"{col}48", f"=SUM({col}7:{col}46)", result)
    for col, numerator in zip("GHI", "DEF"):
        check_formula("Top 40", f"{col}48", f"={numerator}48/C48", totals[numerator] / totals["C"])
    for r, reconciliation in zip(range(23, 26), reconciliations):
        sector = reconciliation["source_sector"]
        check(formulas["Methodik"][f"A{r}"].value == sector, f"Kontrollsparte A{r} falsch")
        check(numeric(formulas["Methodik"][f"C{r}"].value) == reconciliation["branch_sum_million_eur"],
              f"Branchensumme C{r} falsch")
        check_formula("Methodik", f"B{r}", f"=SUMIF('Einzelgesellschaften'!$C$5:$C$330,A{r},'Einzelgesellschaften'!$G$5:$G$330)",
                      reconciliation["entity_sum_million_eur"])
        check_formula("Methodik", f"D{r}", f"=B{r}-C{r}", reconciliation["difference_million_eur"])
        number_count = sum(row["source_sector"] == sector and numeric(row["raw_value"]) is not None for row in rows)
        check_formula("Methodik", f"E{r}", f'=COUNTIFS(\'Einzelgesellschaften\'!$C$5:$C$330,A{r},\'Einzelgesellschaften\'!$D$5:$D$330,">=0")', Decimal(number_count))
    all_formulas = {(s.title, c.coordinate) for s in formulas for row in s.iter_rows() for c in row if c.data_type == "f"}
    check(all_formulas == checked_formulas, "Ungeprüfte oder fehlende Formeln")
    check(workbook_path.read_bytes() == raw, "Eingangsdatei wurde während der Prüfung verändert")
    metadata = {key: cached["Methodik"][cell].value for key, cell in {
        "geography": "B5", "measure": "B6", "grouping": "B7", "sector_limit": "B9", "marker_claim": "B11"}.items()}
    return {
        "schema_version": "ims.ap6.bafin_reference_audit.v1", "evidence_kind": "research_not_selected_dataset",
        "source_workbook": {"filename": workbook_path.name, "bytes": len(raw), "sha256": sha256(raw).hexdigest()},
        "data_year": 2024, "unit": "million_eur", "source_files": provenance,
        "workbook_method_assertions": metadata,
        "audit": {"source_entity_count": len(rows), "group_assertion_count": len(ranked),
                  "formulas_checked": len(checked_formulas), "errors": errors,
                  "arithmetic_and_source_cells_passed": not errors,
                  "group_assignments_independently_verified": False},
        "reconciliation": reconciliations, "displayed_top40_totals": totals,
        "boundary": boundary(ranked), "entities": rows, "groups_by_available_earned_amount": ranked,
        "ap6_gate": {"selection_verified": False, "demo_activation_allowed": False,
                     "open_items": ["German direct business absent; foreign and inward reinsurance included",
                                    "Earned measure differs from proposed written measure",
                                    "German candidate universe incl. EEA not complete",
                                    "Editorial group mapping requires independent historical control evidence",
                                    "Non-life breakdown into motor, property/liability and unmodeled branches absent",
                                    "Common AP6 ranking year and method not accepted"]},
        "marker_resolution": {"source": "AP6-S05", "pdf_page": 2, "printed_page": 1,
                              "dash": "exact_zero", "numeric_zero": "below_source_unit", "other_text": "unknown",
                              "workbook_method_b11_conflicts_with_source": True,
                              "affected_workbook_rows": [r["workbook_row"] for r in rows if r["raw_value"] == "-"],
                              "input_workbook_unchanged": True},
        "precision_note": "Conservative +/- one source unit per nonzero entity; not evidence for geography, consolidation or German rank.",
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--source-dir", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    source_register = json.loads(SOURCE_REGISTER.read_text(encoding="utf-8"))
    inputs = {args.workbook.resolve(), SOURCE_REGISTER.resolve()}
    inputs.update((args.source_dir / Path(s["local_download"]).name).resolve()
                  for s in source_register["sources"] if s["id"] in {v[0] for v in LAYOUT.values()})
    if args.out.resolve() in inputs:
        parser.error("Die Ausgabe darf keine Eingangsdatei überschreiben.")
    result = audit(args.workbook, args.source_dir)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=json_value) + "\n", encoding="utf-8")
    print(json.dumps({"audit": result["audit"], "boundary": result["boundary"], "ap6_gate": result["ap6_gate"]}, ensure_ascii=False, default=json_value))
    # Passing arithmetic is deliberately separate from the still-open AP6 gate.
    return 0 if result["audit"]["arithmetic_and_source_cells_passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
